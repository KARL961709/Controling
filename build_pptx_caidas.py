# -*- coding: utf-8 -*-
"""PPTX nativo y editable: Rebank · Caídas (exclusiones) aplicadas a la base (1 slide).
4 tarjetas de familia con bullets + 2 paneles inferiores. Formas y textos nativos."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

GREEN_DEEP=RGBColor(0x1F,0x51,0x30); GREEN_DEEP2=RGBColor(0x26,0x60,0x3A)
GREEN=RGBColor(0x1A,0xA4,0x4C); GREEN_BRIGHT=RGBColor(0x22,0xB2,0x4C)
GREEN_TEAL=RGBColor(0x2F,0x7D,0x70); GREEN_SOFT=RGBColor(0xE8,0xF5,0xEE)
GRAY_TITLE=RGBColor(0x9A,0xA6,0xA0); GRAY_TXT=RGBColor(0x5B,0x6F,0x66)
INK=RGBColor(0x14,0x27,0x1D); WHITE=RGBColor(0xFF,0xFF,0xFF)
CARD_LINE=RGBColor(0xE7,0xEE,0xE9); AMBER=RGBColor(0xE7,0x9A,0x1E)
AMB_BG=RGBColor(0xFD,0xF3,0xE0); AMB_INK=RGBColor(0x9A,0x6A,0x12); AMB_LINE=RGBColor(0xF2,0xE0,0xBD)
CODE_BG=RGBColor(0xEE,0xF4,0xF1); CODE_INK=RGBColor(0x26,0x60,0x3A)
FOOT_LINE=RGBColor(0xE3,0xEC,0xE7); FOOT_INK=RGBColor(0x8A,0x9B,0x93)
FONT="Segoe UI"; MONO="Consolas"

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
s=prs.slides.add_slide(prs.slide_layouts[6])

def rrect(x,y,w,h,fill,line=None,radius=0.12,lw=1.0):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    try: sh.adjustments[0]=max(0.0,min(0.5,radius/min(w,h)))
    except Exception: pass
    sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    sh.shadow.inherit=False; return sh
def rect(x,y,w,h,fill):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.line.fill.background(); sh.shadow.inherit=False; return sh
def oval(x,y,w,h,fill):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.line.fill.background(); sh.shadow.inherit=False; return sh
def hline(x1,x2,y,color,w):
    ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y),Inches(x2),Inches(y))
    ln.line.color.rgb=color; ln.line.width=Pt(w); ln.shadow.inherit=False; return ln
def tb(x,y,w,h,runs,size=12,color=INK,bold=False,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,ls=1.05,font=FONT):
    bx=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=bx.text_frame; tf.word_wrap=True
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0;tf.vertical_anchor=anchor
    p=tf.paragraphs[0]; p.alignment=align
    if ls:p.line_spacing=ls
    if isinstance(runs,str):runs=[(runs,{})]
    for t,st in runs:
        r=p.add_run();r.text=t;r.font.size=Pt(st.get('sz',size));r.font.bold=st.get('b',bold)
        r.font.name=st.get('f',font);r.font.color.rgb=st.get('c',color)
    return bx
def bullets(x,y,w,h,items,dotcol,size=9.6,gap=0.10):
    bx=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=bx.text_frame; tf.word_wrap=True
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0;tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,runs in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.line_spacing=1.14; p.space_after=Pt(gap*72); p.space_before=Pt(0)
        rb=p.add_run(); rb.text="•  "; rb.font.size=Pt(size); rb.font.bold=True; rb.font.name=FONT; rb.font.color.rgb=dotcol
        for t,st in runs:
            r=p.add_run(); r.text=t; r.font.size=Pt(st.get('sz',size)); r.font.bold=st.get('b',False)
            r.font.name=FONT; r.font.color.rgb=st.get('c',INK)
    return bx

B=lambda c=GREEN_DEEP:{'b':True,'c':c}; MUT={'c':GRAY_TXT,'sz':8.6}

# ---------- Encabezado ----------
tb(0.42,0.28,10.6,0.44,[("Rebank ",{'sz':20,'b':True,'c':GREEN_DEEP}),("| Caídas (exclusiones) aplicadas a la base",{'sz':20,'b':True,'c':GREEN_BRIGHT})],anchor=MSO_ANCHOR.MIDDLE)
tb(0.42,0.74,10.2,0.30,[("Marcas que ",{'sz':10,'c':GRAY_TXT}),("sacan al cliente de la base elegible",{'sz':10,'b':True,'c':GREEN_DEEP}),
   ("  — se conserva a quien ",{'sz':10,'c':GRAY_TXT}),("no tiene ninguna",{'sz':10,'b':True,'c':GREEN_DEEP}),
   (" de estas caídas  ",{'sz':10,'c':GRAY_TXT}),("COALESCE(flg,0) <> 1",{'sz':9,'f':MONO,'c':CODE_INK})])
# badges
rrect(10.62,0.30,1.30,0.36,GREEN_SOFT,radius=0.18)
tb(10.62,0.30,1.30,0.36,[("28",{'sz':13,'b':True,'c':GREEN_DEEP}),(" activas",{'sz':9,'b':True,'c':GREEN_DEEP})],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
rrect(12.00,0.30,0.91,0.36,AMB_BG,radius=0.18)
tb(12.00,0.30,0.91,0.36,[("2",{'sz':13,'b':True,'c':AMB_INK}),(" desact.",{'sz':8.5,'b':True,'c':AMB_INK})],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)

# ---------- 4 tarjetas ----------
COLS=[
 ("⚠","Fraude, listas y clientes no elegibles",8,GREEN_DEEP,[
   [("En ",{}),("base de fraude",B())],
   [("Cliente con ",{}),("marca FEVE",B())],
   [("En lista de ",{}),("clientes especiales",B())],
   [("En base de ",{}),("recovery / cobranza",B())],
   [("Empresa empleadora mala",B()),(" (dudoso / pérdida)",{})],
   [("Empresa no deseada",B())],
   [("Empresa con ",{}),("marca FEVE",B())],
   [("Fallecido",B()),("  (filtro aparte)",MUT)],
 ]),
 ("◆","Riesgo y comportamiento crediticio",9,GREEN_TEAL,[
   [("Con ",{}),("castigo en IBK",B()),(" (últ. 24 meses)",{})],
   [("Con ",{}),("proceso judicial",B()),(" en el SSFF (24 m)",{})],
   [("Clasificación ",{}),("mayor a CPP",B()),(" (buró aprobado)",{})],
   [("Clasificación ",{}),("RCC mayor a Normal",B())],
   [("Más de ",{}),("8 días de mora",B()),(" en IBK",{})],
   [("Comportamiento en IBK ",{}),("distinto de Normal",B())],
   [("Refinanciado",B()),(" en el RCC (últ. 6 meses)",{})],
   [("Refinanciado",B()),(" con buró",{})],
   [("Buró ",{}),("mayor a CPP",B()),(" en los últ. 5 meses",{})],
 ]),
 ("▼","Endeudamiento y capacidad de pago",5,GREEN_BRIGHT,[
   [("Sobreendeudado",B())],
   [("7 o más entidades",B())],
   [("Marca de ",{}),("entidades con saldo",B())],
   [("Sin deuda / ingreso disponible",B())],
   [("Restricción de ",{}),("venta cruzada",B()),(" (PP)",{})],
 ]),
 ("●","Perfil, política y coyuntura",7,GREEN_DEEP2,[
   [("Edad",B()),(" fuera de rango",{})],
   [("Cliente ",{}),("banca especial (BPE)",B())],
   [("Cliente de ",{}),("convenio",B())],
   [("Ya tiene tarjeta de crédito",B())],
   [("Cliente ",{}),("ISR con buró",B())],
   [("Cliente ",{}),("ISR con saldo hipotecario",B()),(" IBK",{})],
   [("Caída por ",{}),("coyuntura",B()),(" (Fenómeno del Niño) – TC",{})],
 ]),
]
L=0.42; gap=0.16; N=4
CW=(13.333-2*L-(N-1)*gap)/N
CY=1.16; CH=3.94; HH=0.62
for i,(ic,title,cnt,col,items) in enumerate(COLS):
    x=L+i*(CW+gap)
    rrect(x,CY,CW,CH,WHITE,line=CARD_LINE,radius=0.11,lw=1.0)
    rrect(x,CY,CW,HH,col,radius=0.11); rect(x,CY+HH-0.11,CW,0.11,col)
    rrect(x+0.12,CY+0.19,0.26,0.26,RGBColor(0x4A,0x7A,0x62),radius=0.06)
    tb(x+0.12,CY+0.19,0.26,0.26,[(ic,{'sz':12,'b':True,'c':WHITE})],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    tb(x+0.46,CY,CW-1.0,HH,[(title,{'sz':12,'b':True,'c':WHITE})],anchor=MSO_ANCHOR.MIDDLE)
    rrect(x+CW-0.56,CY+0.17,0.42,0.30,RGBColor(0x3E,0x74,0x5A),radius=0.06)
    tb(x+CW-0.56,CY+0.17,0.42,0.30,[(str(cnt),{'sz':13,'b':True,'c':WHITE})],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    dot=col
    bullets(x+0.14,CY+HH+0.12,CW-0.28,CH-HH-0.22,items,dot,size=9.6,gap=0.13)

# ---------- 2 paneles inferiores ----------
BY=5.34; BH=1.42
# panel regla
pw1=7.20
rrect(L,BY,pw1,BH,WHITE,line=CARD_LINE,radius=0.11,lw=1.0)
rrect(L,BY,pw1,0.36,GREEN_DEEP,radius=0.11); rect(L,BY+0.25,pw1,0.11,GREEN_DEEP)
tb(L+0.14,BY,pw1-0.2,0.36,[("Cómo se lee la regla",{'sz':11,'b':True,'c':WHITE})],anchor=MSO_ANCHOR.MIDDLE)
# texto con code inline (varios runs)
bx=s.shapes.add_textbox(Inches(L+0.16),Inches(BY+0.48),Inches(pw1-0.32),Inches(BH-0.56))
tf=bx.text_frame; tf.word_wrap=True
tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
p=tf.paragraphs[0]; p.line_spacing=1.35
def run(p,t,sz=11,b=False,c=INK,f=FONT):
    r=p.add_run(); r.text=t; r.font.size=Pt(sz); r.font.bold=b; r.font.name=f; r.font.color.rgb=c
run(p,"Cada condición "); run(p,"COALESCE(flg_x, 0) <> 1",10,False,CODE_INK,MONO); run(p," ")
run(p,"mantiene",11,True,GREEN_DEEP); run(p," a los clientes cuya marca "); run(p,"no está prendida",11,True,GREEN_DEEP)
run(p,": si el cliente tiene la caída, "); run(p,"se excluye",11,True,GREEN_DEEP)
run(p,". Las marcas se arman pivoteando la tabla "); run(p,"caidas",10,False,CODE_INK,MONO)
run(p," (una fila por "); run(p,"descripcion",10,False,CODE_INK,MONO); run(p,") con "); run(p,"MAX(CASE…)",10,False,CODE_INK,MONO); run(p," por cliente.")

# panel desactivadas
px=L+pw1+0.20; pw2=12.91-px
rrect(px,BY,pw2,BH,WHITE,line=CARD_LINE,radius=0.11,lw=1.0)
rrect(px,BY,pw2,0.36,AMBER,radius=0.11); rect(px,BY+0.25,pw2,0.11,AMBER)
tb(px+0.14,BY,pw2-0.2,0.36,[("Desactivadas hoy (comentadas)",{'sz':11,'b':True,'c':WHITE})],anchor=MSO_ANCHOR.MIDDLE)
tb(px+0.16,BY+0.48,pw2-0.32,0.28,[("Estas dos caídas ",{'sz':11,'c':INK}),("no se están aplicando",{'sz':11,'b':True,'c':AMB_INK}),(" en esta corrida:",{'sz':11,'c':INK})])
chy=BY+0.90
for txt in ["Buró bajo en Amazonas / Loreto","Oferta o línea menor al mínimo"]:
    wc=0.30+len(txt)*0.058
    rrect(px+0.16 if txt.startswith("Buró") else px+0.16+ (0.30+len("Buró bajo en Amazonas / Loreto")*0.058)+0.12, chy, wc, 0.30, AMB_BG, line=AMB_LINE, radius=0.10, lw=0.75)
    cx0=px+0.16 if txt.startswith("Buró") else px+0.16+(0.30+len("Buró bajo en Amazonas / Loreto")*0.058)+0.12
    tb(cx0,chy,wc,0.30,[("— "+txt,{'sz':9.6,'b':True,'c':AMB_INK})],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)

# ---------- Pie ----------
hline(0.42,12.91,7.02,FOOT_LINE,1.0)
tb(0.42,7.10,5.6,0.32,[("Base elegible: ",{'sz':7.7,'b':True,'c':GREEN_DEEP2}),("resultado tras aplicar las 28 caídas activas y excluir fallecidos.",{'sz':7.7,'c':FOOT_INK})])
tb(6.3,7.10,6.6,0.32,[("SSFF:",{'sz':7.7,'b':True,'c':GREEN_DEEP2}),(" sistema financiero · ",{'sz':7.7,'c':FOOT_INK}),
   ("RCC:",{'sz':7.7,'b':True,'c':GREEN_DEEP2}),(" reporte crediticio · ",{'sz':7.7,'c':FOOT_INK}),
   ("CPP:",{'sz':7.7,'b':True,'c':GREEN_DEEP2}),(" con problemas potenciales · ",{'sz':7.7,'c':FOOT_INK}),
   ("FEVE:",{'sz':7.7,'b':True,'c':GREEN_DEEP2}),(" marca de alerta comercial.",{'sz':7.7,'c':FOOT_INK})])

prs.save("presentacion_caidas.pptx")
print("OK presentacion_caidas.pptx")
