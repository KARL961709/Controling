# -*- coding: utf-8 -*-
"""PPTX nativo y editable de la Ruta de construcción de árboles (1 slide).
Ruta arriba (6 estaciones, Árbol 1 -> Árbol 2 conectados) + tarjeta de bullets
bajo cada estación. Formas nativas: ovals + rects + textos (nada de imágenes)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

GREEN_DEEP=RGBColor(0x1F,0x51,0x30); GREEN_DEEP2=RGBColor(0x26,0x60,0x3A)
GREEN=RGBColor(0x1A,0xA4,0x4C); GREEN_BRIGHT=RGBColor(0x1A,0x98,0x50)
GREEN_TEAL=RGBColor(0x2F,0x7D,0x70); GREEN_SOFT=RGBColor(0xE8,0xF5,0xEE)
GRAY_TITLE=RGBColor(0x9A,0xA6,0xA0); GRAY_TXT=RGBColor(0x5B,0x6F,0x66)
INK=RGBColor(0x14,0x27,0x1D); WHITE=RGBColor(0xFF,0xFF,0xFF)
CARD_LINE=RGBColor(0xE7,0xEE,0xE9); LINE_SOFT=RGBColor(0xCF,0xE0,0xD8)
HI_BG=RGBColor(0xFF,0xF4,0xE0); HI_INK=RGBColor(0x9A,0x6A,0x12)
CHIP_BG=RGBColor(0xE9,0xEE,0xFB); CHIP_INK=RGBColor(0x34,0x52,0x9C); CHIP_LINE=RGBColor(0xD4,0xDD,0xF6)
FOOT_LINE=RGBColor(0xE3,0xEC,0xE7); FOOT_INK=RGBColor(0x8A,0x9B,0x93)
FONT="Segoe UI"

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
s=prs.slides.add_slide(prs.slide_layouts[6])

def rrect(x,y,w,h,fill,line=None,radius=0.12,lw=1.0):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    try: sh.adjustments[0]=max(0.0,min(0.5,radius/min(w,h)))
    except Exception: pass
    sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    sh.shadow.inherit=False; return sh
def rect(x,y,w,h,fill,line=None,lw=0.75):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    sh.shadow.inherit=False; return sh
def oval(x,y,w,h,fill,line=None,lw=1.0):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    sh.shadow.inherit=False; return sh
def hline(x1,x2,y,color,w):
    ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y),Inches(x2),Inches(y))
    ln.line.color.rgb=color; ln.line.width=Pt(w); ln.shadow.inherit=False; return ln
def tb(x,y,w,h,runs,size=12,color=INK,bold=False,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,ls=1.0):
    bx=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=bx.text_frame; tf.word_wrap=True
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0;tf.vertical_anchor=anchor
    p=tf.paragraphs[0]; p.alignment=align
    if ls:p.line_spacing=ls
    if isinstance(runs,str):runs=[(runs,{})]
    for t,st in runs:
        r=p.add_run();r.text=t;r.font.size=Pt(st.get('sz',size));r.font.bold=st.get('b',bold)
        r.font.name=FONT;r.font.color.rgb=st.get('c',color)
    return bx
def bullets(x,y,w,h,items,size=8,gap=0.12):
    """items: lista de bullets; cada bullet = lista de runs (t,{sz,b,c})."""
    bx=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=bx.text_frame; tf.word_wrap=True
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0;tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,runs in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.line_spacing=1.12; p.space_after=Pt(gap*72); p.space_before=Pt(0)
        rb=p.add_run(); rb.text="•  "; rb.font.size=Pt(size); rb.font.bold=True
        rb.font.name=FONT; rb.font.color.rgb=GREEN
        for t,st in runs:
            r=p.add_run(); r.text=t; r.font.size=Pt(st.get('sz',size)); r.font.bold=st.get('b',False)
            r.font.name=FONT; r.font.color.rgb=st.get('c',INK)
    return bx

B={'b':True,'c':GREEN_DEEP}   # negrita verde
BI={'b':True}                 # negrita ink

# ---------- Encabezado ----------
tb(0.42,0.30,10.5,0.42,[("Ruta de construcción ",{'sz':19,'b':True,'c':GREEN_DEEP}),
                        ("| Score Rebank · el Árbol 1 alimenta al Árbol 2",{'sz':19,'b':True,'c':GRAY_TITLE})],anchor=MSO_ANCHOR.MIDDLE)
tb(0.42,0.74,10.6,0.28,[("Un solo recorrido: del score a los buckets finales · el ",{'sz':10,'c':GRAY_TXT}),
                        ("resultado del Árbol 1 entra como insumo del Árbol 2",{'sz':10,'b':True,'c':GREEN_DEEP})])
# badge
rrect(11.05,0.34,1.86,0.34,GREEN_SOFT,radius=0.17)
oval(11.22,0.45,0.11,0.11,GREEN)
tb(11.40,0.34,1.5,0.34,[("24 meses y 5 años",{'sz':9,'b':True,'c':GREEN_DEEP})],anchor=MSO_ANCHOR.MIDDLE)

# ---------- Geometría de columnas ----------
mL=0.42; g=0.13; N=6
cardW=(13.333-2*mL-(N-1)*g)/N            # ~1.973
xs=[mL+i*(cardW+g) for i in range(N)]    # left de cada tarjeta
cx=[x+cardW/2 for x in xs]               # centro (estación)
cy=1.66; r=0.215
colhead=[GREEN_DEEP,GREEN_TEAL,GREEN_TEAL,GREEN_BRIGHT,GREEN_BRIGHT,GREEN_DEEP]

# ---------- Ruta (línea + tramo enfatizado + estaciones) ----------
hline(cx[0],cx[5],cy,LINE_SOFT,4.5)
# tramo 3->4 enfatizado
hline(cx[2],cx[3],cy,GREEN,5.5)
# flecha hacia estación 4
ar=cx[3]-0.055
tribase=[(ar-0.14,cy-0.11),(ar+0.02,cy),(ar-0.14,cy+0.11)]
from pptx.enum.shapes import MSO_SHAPE as _S
fb=s.shapes.build_freeform(int(tribase[0][0]*914400),int(tribase[0][1]*914400),scale=1)
fb.add_line_segments([(int(px*914400),int(py*914400)) for px,py in tribase[1:]],close=True)
tri=fb.convert_to_shape(); tri.fill.solid(); tri.fill.fore_color.rgb=GREEN; tri.line.fill.background(); tri.shadow.inherit=False
# texto del tramo enfatizado
midx=(cx[2]+cx[3])/2
tb(midx-1.1,cy-0.55,2.2,0.2,[("el bucket del Árbol 1",{'sz':9.5,'b':True,'c':RGBColor(0x13,0x7A,0x4A)})],align=PP_ALIGN.CENTER)
tb(midx-1.1,cy+0.30,2.2,0.2,[("entra al Árbol 2",{'sz':8.5,'c':GRAY_TXT})],align=PP_ALIGN.CENTER)

# brackets de grupo
def bracket(x1,x2,ytop,color,label):
    tb((x1+x2)/2-1.3,ytop-0.02,2.6,0.2,[(label,{'sz':9.5,'b':True,'c':color})],align=PP_ALIGN.CENTER)
    yb=ytop+0.20
    hline(x1,x2,yb,color,1.4)
    ln1=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(yb),Inches(x1),Inches(yb+0.08))
    ln1.line.color.rgb=color; ln1.line.width=Pt(1.4); ln1.shadow.inherit=False
    ln2=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x2),Inches(yb),Inches(x2),Inches(yb+0.08))
    ln2.line.color.rgb=color; ln2.line.width=Pt(1.4); ln2.shadow.inherit=False
bracket(cx[1],cx[2],1.06,GREEN_TEAL,"ÁRBOL 1 · segmentación base")
bracket(cx[3],cx[4],1.06,GREEN_BRIGHT,"ÁRBOL 2 · refinamiento")

# estaciones (círculos numerados)
for i in range(N):
    oval(cx[i]-r,cy-r,2*r,2*r,colhead[i],line=WHITE,lw=2.5)
    tb(cx[i]-r,cy-r,2*r,2*r,[(str(i+1),{'sz':14,'b':True,'c':WHITE})],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)

# ---------- Tarjetas ----------
cardTop=2.12; cardH=4.45; headH=0.40
titles=["Score Rebank","Árbol 1 · segmentación","Buckets del Árbol 1",
        "Árbol 2 · refinamiento","Buckets finales","Validación"]
CONTENT=[
 [[("Se parte de la ",{}),("base de desarrollo",B),(" (clientes con resultado conocido de riesgo).",{})],
  [("El score se ",{}),("agrupa en 6 tramos",B),(" según su nivel de riesgo.",{})],
  [("Resultado: ",{}),("categoría de Score Rebank",B),(", de ",{}),("1 (mejor)",B),(" a ",{}),("6 (peor)",B),(".",{})]],
 [[("Entra: ",{}),("variables del cliente",B),(" — segmento comercial, ranking de ingreso y comportamiento de castigo (deuda, entidades, meses desde el castigo).",{})],
  [("Cada variable se ",{}),("agrupa en tramos ordenados por riesgo",B),("; sin dato = mayor riesgo.",{})],
  [("Un ",{}),("árbol de decisión",B),(" que respeta la dirección de riesgo de cada variable (se prueban profundidades 4 a 7).",{})],
  [("Se conserva ",{}),("solo las variables que aportan",B),(".",{})],
  [("Salida: ",{}),("estrategias",B),(" (hojas) ordenadas por riesgo.",{})]],
 [[("Las estrategias se ",{}),("agrupan en 7 buckets",B),(", de menor a mayor riesgo.",{})],
  [("24 meses:",BI),(" 21 estrategias → 7 buckets.",{})],
  [("5 años:",BI),(" 15 estrategias → 7 buckets.",{})]],   # el bullet resaltado va aparte
 [[("Entra: el ",{}),("bucket del Árbol 1",B),(" + el ",{}),("saldo pasivo promedio (3 meses)",B),(" + la ",{}),("categoría de Score Rebank",B),(".",{})],
  [("Mismo procedimiento",BI),(": agrupar variables por riesgo → árbol de decisión monótono → conservar solo lo que aporta.",{})],
  [("De las variables de ",{}),("saldo pasivo",B),(" se deja ",{}),("una sola",B),(" (la más importante).",{})],
  [("Salida: ",{}),("estrategias más finas",B),(".",{})]],
 [[("Las estrategias se ",{}),("agrupan en 12 buckets finales",B),(" ordenados por riesgo.",{})],
  [("24 meses:",BI),(" 30 estrategias → 12 · ",{}),("5 años:",BI),(" 16 → 12.",{})],
  [("Métodos de ",{}),("ordenamiento",B),(" probados:",{})]],   # chips aparte
 [[("Base de campañas:",BI),(" clientes que hoy se envían en los pilotos.",{})],
  [("Base inicial:",BI),(" la base con la que se construyen los árboles.",{})],
  [("Base fuera de campaña:",BI),(" clientes que no están en los pilotos actuales.",{})],
  [("En cada una se mide la ",{}),("distribución de los buckets",B),(".",{})]],
]

for i in range(N):
    x=xs[i]
    rrect(x,cardTop,cardW,cardH,WHITE,line=CARD_LINE,radius=0.11,lw=1.0)
    rrect(x,cardTop,cardW,headH,colhead[i],radius=0.11)
    rect(x,cardTop+headH-0.11,cardW,0.11,colhead[i])  # tapa las esquinas inferiores redondeadas del header
    # número + título en el header
    rrect(x+0.10,cardTop+0.09,0.22,0.22,RGBColor(0x4A,0x7A,0x62) if colhead[i]==GREEN_DEEP else RGBColor(0x5F,0x9E,0x83),radius=0.06)
    tb(x+0.10,cardTop+0.09,0.22,0.22,[(str(i+1),{'sz':9,'b':True,'c':WHITE})],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    tb(x+0.38,cardTop,cardW-0.42,headH,[(titles[i],{'sz':9.3,'b':True,'c':WHITE})],anchor=MSO_ANCHOR.MIDDLE)
    by=cardTop+headH+0.10
    bullets(x+0.12,by,cardW-0.24,cardH-headH-0.2,CONTENT[i],size=8,gap=0.055)

# extras: bullet resaltado tarjeta 3
hy=cardTop+headH+0.10+3*0.34+0.02
rrect(xs[2]+0.10,hy,cardW-0.20,0.40,HI_BG,radius=0.06)
tb(xs[2]+0.20,hy,cardW-0.34,0.40,[("Este bucket es el insumo del Árbol 2.",{'sz':8,'b':True,'c':HI_INK})],anchor=MSO_ANCHOR.MIDDLE)

# chips de métodos tarjeta 5
chips=["Cascada","ChiMerge","OptBin","Shrinkage","MDLP"]
chx=xs[4]+0.12; chy=cardTop+headH+0.10+3*0.34+0.04; chh=0.20
cxx=chx
for j,ch in enumerate(chips):
    wch=0.14+len(ch)*0.052
    if cxx+wch>xs[4]+cardW-0.12:
        cxx=chx; chy+=chh+0.06
    rrect(cxx,chy,wch,chh,CHIP_BG,line=CHIP_LINE,radius=0.09,lw=0.75)
    tb(cxx,chy,wch,chh,[(ch,{'sz':7.5,'b':True,'c':CHIP_INK})],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    cxx+=wch+0.06

# ---------- Pie ----------
hline(0.42,12.91,6.78,FOOT_LINE,1.0)
tb(0.42,6.86,7.0,0.5,[("Objetivo del árbol: ",{'sz':7.5,'b':True,'c':GREEN_DEEP2}),
   ("ordenar por score (mayor score → menor riesgo) · las variables entran con su dirección de riesgo esperada.",{'sz':7.5,'c':FOOT_INK})])
tb(7.6,6.86,5.3,0.5,[("Árbol 1 → Árbol 2: ",{'sz':7.5,'b':True,'c':GREEN_DEEP2}),
   ("el bucket del primer árbol es una entrada del segundo · el saldo pasivo solo aparece en el Árbol 2.",{'sz':7.5,'c':FOOT_INK})])

prs.save("presentacion_ruta_construccion.pptx")
print("OK presentacion_ruta_construccion.pptx")
