# -*- coding: utf-8 -*-
"""Auditoría geométrica del PPTX: estima el alto real de cada tabla según el
ajuste de texto (word-wrap) y verifica que ninguna forma se salga de la
diapositiva, que las tablas quepan en su tarjeta y que nada invada el pie."""
import sys
from pptx import Presentation
from pptx.util import Emu

EMU = 914400.0  # por pulgada
SW, SH = 13.333, 7.5
FOOT_Y = 7.02

def inch(v): return v / EMU

def chars_per_line(width_in, size_pt, mono=False):
    cpi = (120.0 if mono else 144.0) / size_pt   # chars por pulgada (aprox)
    return max(1, width_in * cpi)

def line_h(size_pt):
    return size_pt * 1.2 / 72.0

def text_len(cell):
    return sum(len(r.text) for p in cell.text_frame.paragraphs for r in p.runs)

def max_size(cell, default=10):
    sizes = [r.font.size.pt for p in cell.text_frame.paragraphs for r in p.runs if r.font.size]
    return max(sizes) if sizes else default

def is_mono(cell):
    for p in cell.text_frame.paragraphs:
        for r in p.runs:
            if r.font.name and 'Consol' in r.font.name:
                return True
    return False

def est_table_height(tbl_graphic):
    """Estima alto renderizado de la tabla usando los márgenes reales de cada celda."""
    tbl = tbl_graphic.table
    col_w = [inch(c.width) for c in tbl.columns]
    total = 0.0
    detail = []
    for ri in range(len(tbl.rows)):
        needed = inch(tbl.rows[ri].height)   # alto mínimo fijado
        for ci in range(len(col_w)):
            cell = tbl.cell(ri, ci)
            n = text_len(cell)
            if n == 0:
                continue
            sz = max_size(cell)
            mt = cell.margin_top / EMU if cell.margin_top is not None else 0.05
            mb = cell.margin_bottom / EMU if cell.margin_bottom is not None else 0.05
            ww = cell.text_frame.word_wrap
            if ww is False:
                lines = 1
            else:
                cpl = chars_per_line(col_w[ci] - (mt + mb) - 0.06, sz, is_mono(cell))
                lines = max(1, -(-n // max(1, int(cpl))))
            h = lines * line_h(sz) + mt + mb + 0.025  # +buffer ascendente/descendente
            needed = max(needed, h)
        total += needed
        detail.append(round(needed, 3))
    return total, detail

def is_white_card(sh):
    try:
        f = sh.fill
        if str(f.fore_color.rgb) == 'FFFFFF' and inch(sh.width) > 1.5 and inch(sh.height) > 0.8:
            return True
    except Exception:
        return False
    return False

def is_textbox(sh):
    try:
        return sh.has_text_frame and not sh.has_table and sh.shape_type is not None and 'TEXT_BOX' in str(sh.shape_type)
    except Exception:
        return False

prs = Presentation(sys.argv[1] if len(sys.argv) > 1 else "presentacion_metodologia_arboles.pptx")
problems = []
for si, slide in enumerate(prs.slides, 1):
    print(f"\n===== SLIDE {si} =====")
    tables = []
    cards = []
    boxes = []
    for sh in slide.shapes:
        if is_white_card(sh):
            cards.append((inch(sh.left), inch(sh.top), inch(sh.left)+inch(sh.width), inch(sh.top)+inch(sh.height)))
        if is_textbox(sh):
            boxes.append((inch(sh.left), inch(sh.top), inch(sh.left)+inch(sh.width), inch(sh.top)+inch(sh.height)))
    for sh in slide.shapes:
        L, T = inch(sh.left), inch(sh.top)
        W, H = inch(sh.width), inch(sh.height)
        R, B = L + W, T + H
        # fuera de límites
        if L < -0.02 or T < -0.02 or R > SW + 0.02 or B > SH + 0.02:
            problems.append(f"S{si}: forma fuera de límites '{sh.shape_type}' L={L:.2f} T={T:.2f} R={R:.2f} B={B:.2f}")
        if sh.has_table:
            th, det = est_table_height(sh.graphic_frame if hasattr(sh, 'graphic_frame') else sh)
            tables.append((sh, T, th))
            alloc = H
            bottom = T + th
            # tarjeta contenedora = la menor que contiene el (L,T) de la tabla
            host = None
            for (cl, ct, cr, cb) in cards:
                if cl - 0.05 <= L and R <= cr + 0.05 and ct - 0.05 <= T and cb + 0.05 >= T:
                    if host is None or (cb - ct) < (host[3] - host[1]):
                        host = (cl, ct, cr, cb)
            cardb = host[3] if host else None
            tag = f" card_bottom={cardb:.2f}" if cardb else ""
            print(f"  TABLA top={T:.2f} alloc_h={alloc:.2f} est_h={th:.2f} bottom_est={bottom:.2f}{tag}  filas={det}")
            if bottom > FOOT_Y - 0.02:
                problems.append(f"S{si}: tabla invade el pie (bottom_est={bottom:.2f} > {FOOT_Y})")
            if cardb and bottom > cardb - 0.04:
                problems.append(f"S{si}: tabla se sale de su tarjeta (bottom_est={bottom:.2f} > card_bottom={cardb:.2f})")
    # textos que invaden la región estimada de una tabla
    for (sh, T, th) in tables:
        tL, tR = inch(sh.left), inch(sh.left)+inch(sh.width)
        tT, tB = T, T + th
        for (bl, bt, br, bb) in boxes:
            ox = not (br <= tL + 0.05 or bl >= tR - 0.05)
            oy = not (bb <= tT + 0.03 or bt >= tB - 0.03)
            if ox and oy:
                problems.append(f"S{si}: un texto invade la región de una tabla (texto top={bt:.2f} vs tabla {tT:.2f}-{tB:.2f})")

    # chequear que tablas no se solapen entre sí (verticalmente, misma columna)
    for i in range(len(tables)):
        for j in range(i+1, len(tables)):
            a, at, ah = tables[i]; b, bt, bh = tables[j]
            aL, aR = inch(a.left), inch(a.left)+inch(a.width)
            bL, bR = inch(b.left), inch(b.left)+inch(b.width)
            overlapX = not (aR <= bL + 0.02 or bR <= aL + 0.02)
            ab, bb = at+ah, bt+bh
            overlapY = not (ab <= bt + 0.02 or bb <= at + 0.02)
            if overlapX and overlapY:
                problems.append(f"S{si}: dos tablas se solapan (top {at:.2f}/{bt:.2f})")

print("\n===== RESULTADO AUDITORÍA =====")
if problems:
    for p in problems:
        print("  ✗", p)
    sys.exit(1)
else:
    print("  ✓ Sin formas fuera de límites, sin tablas invadiendo el pie ni solapadas.")
