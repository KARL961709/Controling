# -*- coding: utf-8 -*-
"""PPTX nativo y editable del deck de buckets (4 slides) — replica
presentacion_buckets_v1v2v3.html (layout 2x2 + tabla V1 pivoteada)."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

GREEN_DEEP=RGBColor(0x1F,0x51,0x30); GREEN_DEEP2=RGBColor(0x26,0x60,0x3A)
GREEN=RGBColor(0x1A,0xA4,0x4C); GREEN_BRIGHT=RGBColor(0x1A,0x98,0x50)
GREEN_TEAL=RGBColor(0x3A,0x7D,0x5C); GREEN_SOFT=RGBColor(0xE8,0xF5,0xEE)
TH_BG=RGBColor(0x0A,0x5A,0x44); GRAY_TITLE=RGBColor(0x9A,0xA6,0xA0)
GRAY_TXT=RGBColor(0x5B,0x6F,0x66); INK=RGBColor(0x14,0x27,0x1D)
MONO_INK=RGBColor(0x23,0x43,0x3C); WHITE=RGBColor(0xFF,0xFF,0xFF)
CARD_LINE=RGBColor(0xE7,0xEE,0xE9); BORDER="DDE7E4"
TOT_BG=RGBColor(0xCF,0xE9,0xE4); TOT_INK=RGBColor(0x08,0x4B,0x44)
RJ_BG=RGBColor(0xEC,0xEF,0xF0); RJ_INK=RGBColor(0x4A,0x5A,0x56)
TOTC_BG=RGBColor(0xDF,0xF0,0xED); TOTC_INK=RGBColor(0x0A,0x5A,0x52)
FONT="Segoe UI"; MONO="Consolas"
C=PP_ALIGN.CENTER; L=PP_ALIGN.LEFT; R=PP_ALIGN.RIGHT

def hsl(h, s=0.60, l=0.78):
    c=(1-abs(2*l-1))*s; x=c*(1-abs((h/60.0)%2-1)); m=l-c/2
    if   h<60:  r,g,b=c,x,0
    elif h<120: r,g,b=x,c,0
    elif h<180: r,g,b=0,c,x
    elif h<240: r,g,b=0,x,c
    elif h<300: r,g,b=x,0,c
    else:       r,g,b=c,0,x
    return RGBColor(int((r+m)*255), int((g+m)*255), int((b+m)*255))

BUCK = [hsl(130-(k)/6.0*130) for k in range(7)]   # B1..B7 verde->rojo

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]

def rrect(s,x,y,w,h,fill,line=None,radius=0.12,lw=1.0):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    try: sh.adjustments[0]=max(0.0,min(0.5,radius/min(w,h)))
    except Exception: pass
    sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    sh.shadow.inherit=False; return sh

def rect(s,x,y,w,h,fill):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.line.fill.background(); sh.shadow.inherit=False; return sh

def tb(s,x,y,w,h,runs,size=12,color=INK,bold=False,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,font=FONT,ls=1.0):
    bx=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=bx.text_frame; tf.word_wrap=True
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0; tf.vertical_anchor=anchor
    p=tf.paragraphs[0]; p.alignment=align
    if ls: p.line_spacing=ls
    if isinstance(runs,str): runs=[(runs,{})]
    for t,st in runs:
        r=p.add_run(); r.text=t; r.font.size=Pt(st.get('sz',size)); r.font.bold=st.get('b',bold)
        r.font.name=st.get('f',font); r.font.color.rgb=st.get('c',color)
    return bx

def set_border(cell,color=BORDER,w=9525):
    tcPr=cell._tc.get_or_add_tcPr()
    for tag in ('a:lnB','a:lnT','a:lnR','a:lnL'):
        for el in tcPr.findall(qn(tag)): tcPr.remove(el)
        ln=tcPr.makeelement(qn(tag),{'w':str(w),'cap':'flat'})
        sf=tcPr.makeelement(qn('a:solidFill'),{}); cl=tcPr.makeelement(qn('a:srgbClr'),{'val':color})
        sf.append(cl); ln.append(sf); tcPr.insert(0,ln)

def cell(c,runs,size=11,color=INK,bold=False,align=PP_ALIGN.CENTER,fill=None,font=FONT,anchor=MSO_ANCHOR.MIDDLE,mgn=0.03,nowrap=False):
    if fill is not None: c.fill.solid(); c.fill.fore_color.rgb=fill
    else: c.fill.background()
    c.vertical_anchor=anchor; c.margin_left=Inches(0.05); c.margin_right=Inches(0.04)
    c.margin_top=Inches(mgn); c.margin_bottom=Inches(mgn)
    tf=c.text_frame; tf.word_wrap=(not nowrap)
    p=tf.paragraphs[0]; p.alignment=align
    if isinstance(runs,str): runs=[(runs,{})]
    first=True
    for t,st in runs:
        r=p.runs[0] if (first and p.runs) else p.add_run(); first=False
        r.text=t; r.font.size=Pt(st.get('sz',size)); r.font.bold=st.get('b',bold)
        r.font.name=st.get('f',font); r.font.color.rgb=st.get('c',color)
    set_border(c)

def card(s,x,y,w,h,title,hcolor=GREEN_DEEP,sub=None,hh=0.34):
    rrect(s,x,y,w,h,WHITE,line=CARD_LINE,radius=0.12)
    rrect(s,x,y,w,hh,hcolor,radius=0.12); rect(s,x,y+hh-0.13,w,0.13,hcolor)
    runs=[(title,{'b':True,'sz':13,'c':WHITE})]
    if sub: runs.append(("   "+sub,{'sz':10,'c':WHITE}))
    tb(s,x+0.15,y,w-0.3,hh,runs,anchor=MSO_ANCHOR.MIDDLE)
    return x+0.14,y+hh+0.07,w-0.28

def header(s,title_runs,subtitle,badge):
    tb(s,0.46,0.26,9.8,0.5,title_runs,size=25,bold=True,anchor=MSO_ANCHOR.MIDDLE)
    tb(s,0.46,0.78,10.0,0.3,subtitle,size=11,color=GRAY_TXT,anchor=MSO_ANCHOR.MIDDLE)
    bw=2.5; rrect(s,13.333-0.46-bw,0.28,bw,0.42,GREEN_SOFT,radius=0.21)
    tb(s,13.333-0.46-bw,0.28,bw,0.42,[("●  ",{'c':GREEN,'b':True,'sz':10}),(badge,{'c':GREEN_DEEP,'b':True,'sz':11})],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)

def footer(s,items):
    rect(s,0.46,7.02,13.333-0.92,0.012,RGBColor(0xE3,0xEC,0xE7))
    runs=[]
    for i,(lead,rest) in enumerate(items):
        if i: runs.append(("      ",{}))
        runs.append((lead,{'b':True,'c':GREEN_DEEP2,'sz':8})); runs.append((rest,{'c':GRAY_TITLE,'sz':8}))
    tb(s,0.46,7.08,13.333-0.92,0.34,runs,size=8,color=GRAY_TITLE)

def add_slide(): return prs.slides.add_slide(BLANK)

def th_row(t,row,headers,size=10.5):
    for j,(htxt,al,sp) in enumerate(headers):
        fill=TH_BG; col=WHITE
        if sp=='rj': fill=RJ_BG; col=RJ_INK
        if sp=='tot': fill=TOTC_BG; col=TOTC_INK
        cell(t.cell(row,j),[(htxt,{'sz':size})],color=col,bold=True,fill=fill,align=al)

def buckcell(c,label,hue_idx,size=10.5,mgn=0.02):
    cell(c,[(label,{'sz':size})],color=INK,bold=True,fill=BUCK[hue_idx],align=PP_ALIGN.CENTER,mgn=mgn)

def setw(t,ws):
    for j,wd in enumerate(ws): t.columns[j].width=Inches(wd)

# ====== datos resumen ======
# Tasa: (Estrat, Tasa, Prob, N, %tot)
TASA_24=[("1, 2, 3","16.89%","31.6%","1,403","3.1%"),("4, 5","20.56%","35.3%","214","0.5%"),
 ("6–10","20.75%","40.6%","32,719","72%"),("11, 12, 13","21.41%","42.4%","4,611","10.1%"),
 ("14–17","27.28%","44.4%","942","2.1%"),("18, 19","26.86%","45.7%","4,318","9.5%"),
 ("20, 21","26.92%","47.5%","1,252","2.8%")]
TASA_24_TOT=("21 est.","21.58%","41.2%","45,459","100%")
# V2/V3: (Cantidad, %, Prob, score)
V2_24=[("19,103","3.3%","33%","670"),("19,102","3.3%","35.4%","646"),("76,066","13%","39.5%","605"),
 ("150,105","25.8%","42.2%","577"),("245,492","42.1%","44.6%","553"),("53,558","9.2%","46.3%","537"),
 ("19,501","3.3%","48.2%","518")]; V2_24_TOT=("582,927","100%","42.9%","570")
V3_24=[("13,317","2.5%","33.7%","663"),("13,673","2.6%","36.3%","636"),("57,746","11%","40.4%","596"),
 ("125,534","24%","43%","569"),("240,860","46%","44.7%","552"),("52,916","10.1%","46.3%","536"),
 ("19,438","3.7%","48.2%","517")]; V3_24_TOT=("523,484","100%","43.6%","563")
# V1 pivote: tipos + filas por bucket (valor por tipo, Total)
V1_24_TIPOS=["TC_CMP_EST_CAST_24_T_CL1_3","TC_CMP_EST_CAST_24_T_CL4","TC_CMP_EST_CAS_24_IC"]
V1_24=[("9,641","675","0","10,316"),("0","6,807","0","6,807"),("12,901","4,106","0","17,007"),
 ("0","21,133","0","21,133"),("0","0","3,930","3,930"),("0","0","472","472"),("0","0","42","42")]
V1_24_TOT=("22,542","32,721","4,444","59,707")

TASA_5=[("1","16.82%","32.1%","1,290","2.8%"),("2, 3","19.75%","35.5%","319","0.7%"),
 ("5, 6, 7","20.75%","38.5%","32,661","71.8%"),("9, 10","18.26%","40.1%","115","0.3%"),
 ("12","21.46%","40.9%","4,562","10%"),("13, 14","26.96%","42.8%","6,496","14.3%"),
 ("15","18.75%","44.6%","16","0%")]
TASA_5_TOT=("15 est.","21.58%","39.1%","45,459","100%")
V2_5=[("12,897","0.8%","32.1%","679"),("51,199","3%","35.6%","644"),("230,088","13.6%","38.8%","612"),
 ("98,755","5.8%","40.1%","599"),("533,913","31.4%","40.9%","591"),("612,572","36.1%","42.9%","571"),
 ("158,251","9.3%","44.6%","553")]; V2_5_TOT=("1,697,675","100%","41.4%","586")
V3_5=[("9,525","0.6%","33.5%","665"),("45,934","2.8%","36.2%","637"),("212,618","12.9%","39.3%","607"),
 ("94,600","5.7%","40.4%","596"),("522,188","31.7%","41%","590"),("607,375","36.8%","42.9%","570"),
 ("157,434","9.5%","44.7%","553")]; V3_5_TOT=("1,649,674","100%","41.6%","583")
V1_5_TIPOS=["TC_CMP_EST_5A_CAST_T","TC_CMP_EST_CAS_5A_IC"]
V1_5=[("2,886","837","3,723"),("1,195","3,881","5,076"),("7,532","6,520","14,052"),
 ("1,198","2,015","3,213"),("0","10,097","10,097"),("0","4,486","4,486"),("0","740","740")]
V1_5_TOT=("12,811","28,576","41,387")

# ====== datos reglas (sin cambios) ======
REGLAS_24=[("B1",0,[("1","segmentacion_gdp_v2 ≤ 2.5 & rk_ing_num > 3,383.5","16.8%","2.8%","1,294"),
  ("2","segmentacion_gdp_v2 (2.5, 3.5] & rk_ing_num > 3,383.5 & meses_desde_primer_castigo > 22.5","16.5%","0.2%","79"),
  ("3","segmentacion_gdp_v2 (2.5, 3.5] & rk_ing_num > 3,383.5 & meses_desde_primer_castigo ≤ 22.5","20.0%","0.1%","30")]),
 ("B2",1,[("4","segmentacion_gdp_v2 (3.5, 4.5] & rk_ing_num > 3,383.5 & meses_desde_primer_castigo > 22.5","19.0%","0.3%","158"),
  ("5","segmentacion_gdp_v2 (3.5, 4.5] & rk_ing_num > 3,383.5 & meses_desde_primer_castigo ≤ 22.5","25.0%","0.1%","56")]),
 ("B3",2,[("6","segmentacion_gdp_v2 ≤ 3.5 & rk_ing_num (3,088.5, 3,383.5]","17.3%","0.7%","301"),
  ("7","segmentacion_gdp_v2 ≤ 3.5 & rk_ing_num (2,746.5, 3,088.5]","20.5%","1.2%","533"),
  ("8","segmentacion_gdp_v2 ≤ 3.5 & rk_ing_num (1,411.5, 2,746.5] & meses_desde_primer_castigo > 22.5","19.1%","4.6%","2,071"),
  ("9","segmentacion_gdp_v2 ≤ 3.5 & rk_ing_num (1,411.5, 2,746.5] & meses_desde_primer_castigo ≤ 22.5","20.1%","1.6%","730"),
  ("10","segmentacion_gdp_v2 ≤ 3.5 & rk_ing_num ≤ 1,411.5","20.9%","64%","29,084")]),
 ("B4",3,[("11","segmentacion_gdp_v2 (3.5, 4.5] & rk_ing_num (2,910.5, 3,383.5]","23.7%","0.2%","97"),
  ("12","segmentacion_gdp_v2 (3.5, 4.5] & rk_ing_num (2,746.5, 2,910.5]","15.7%","0.1%","51"),
  ("13","segmentacion_gdp_v2 (3.5, 4.5] & rk_ing_num ≤ 2,746.5","21.4%","9.8%","4,463")]),
 ("B5",4,[("14","segmentacion_gdp_v2 > 4.5 & rk_ing_num > 2,746.5 & meses_desde_primer_castigo > 22.5","26.4%","0.6%","269"),
  ("15","segmentacion_gdp_v2 > 4.5 & rk_ing_num > 2,746.5 & meses_desde_primer_castigo (9.5, 22.5]","17.6%","0.2%","74"),
  ("16","segmentacion_gdp_v2 > 4.5 & rk_ing_num > 2,746.5 & meses_desde_primer_castigo ≤ 9.5","20.0%","0%","5"),
  ("17","segmentacion_gdp_v2 > 4.5 & rk_ing_num (1,348.5, 2,746.5] & meses_desde_primer_castigo > 22.5","29.0%","1.3%","594")]),
 ("B6",5,[("18","segmentacion_gdp_v2 > 4.5 & rk_ing_num ≤ 1,348.5 & meses_desde_primer_castigo > 22.5","26.8%","9.1%","4,155"),
  ("19","segmentacion_gdp_v2 > 4.5 & rk_ing_num (1,348.5, 2,746.5] & meses_desde_primer_castigo (9.5, 22.5]","28.2%","0.4%","163")]),
 ("B7",6,[("20","segmentacion_gdp_v2 > 4.5 & rk_ing_num ≤ 1,348.5 & meses_desde_primer_castigo (9.5, 22.5]","27.5%","2.5%","1,152"),
  ("21","segmentacion_gdp_v2 > 4.5 & rk_ing_num ≤ 2,746.5 & meses_desde_primer_castigo ≤ 9.5","20.0%","0.2%","100")])]

REGLAS_5=[("B1",0,[("1","segmentacion_gdp_v2 ≤ 2.5 & rk_ing_num > 3,391.5","16.8%","2.8%","1,290")]),
 ("B2",1,[("2","segmentacion_gdp_v2 (2.5, 3.5] & rk_ing_num > 3,391.5 & nro_entidades_castigo ≤ 1.5","17.4%","0.2%","109"),
  ("3","segmentacion_gdp_v2 (3.5, 4.5] & rk_ing_num > 3,391.5 & nro_entidades_castigo ≤ 1.5","21.0%","0.5%","210")]),
 ("B3",2,[("5","segmentacion_gdp_v2 ≤ 3.5 & rk_ing_num ≤ 3,391.5 & nro_entidades_castigo ≤ 1.5 & deuda_cas ≤ 218.33","20.3%","23.8%","10,802"),
  ("6","segmentacion_gdp_v2 ≤ 3.5 & rk_ing_num (3,100.5, 3,391.5] & nro_entidades_castigo ≤ 1.5 & deuda_cas > 218.33","21.6%","0.4%","185"),
  ("7","segmentacion_gdp_v2 ≤ 3.5 & rk_ing_num ≤ 3,100.5 & nro_entidades_castigo ≤ 1.5 & deuda_cas > 218.33","21.0%","47.7%","21,674")]),
 ("B4",3,[("9","segmentacion_gdp_v2 ≤ 3.5 & rk_ing_num ≤ 3,100.5 & nro_entidades_castigo > 1.5","21.0%","0.1%","62"),
  ("10","segmentacion_gdp_v2 (3.5, 4.5] & rk_ing_num (3,100.5, 3,391.5] & nro_entidades_castigo ≤ 1.5","15.1%","0.1%","53")]),
 ("B5",4,[("12","segmentacion_gdp_v2 (3.5, 4.5] & rk_ing_num ≤ 3,100.5","21.5%","10%","4,562")]),
 ("B6",5,[("13","segmentacion_gdp_v2 > 4.5 & nro_entidades_castigo ≤ 1.5 & deuda_cas ≤ 218.33","26.7%","4%","1,818"),
  ("14","segmentacion_gdp_v2 > 4.5 & nro_entidades_castigo ≤ 1.5 & deuda_cas > 218.33","27.1%","10.3%","4,678")]),
 ("B7",6,[("15","segmentacion_gdp_v2 > 4.5 & nro_entidades_castigo > 1.5","18.8%","0%","16")])]


def dist_table(s,cx,cy,cw,ttl,data,tot):
    """Card V2/V3: Buck | Cantidad | % | Prob. | score."""
    ix,iy,iw=card(s,cx,cy,cw,TH,ttl,hcolor=GREEN_BRIGHT)
    t=s.shapes.add_table(9,5,Inches(ix),Inches(iy),Inches(iw),Inches(2.144)).table
    t.first_row=False; t.horz_banding=False
    ws=[0.5, iw-0.5-1.05-1.35-1.4, 1.05, 1.35, 1.4]
    setw(t,ws)
    th_row(t,0,[("Buck.",C,''),("Cantidad",C,''),("%",C,''),("Prob.",C,''),("score",C,'')])
    for i,(cn,pc,pj,sc) in enumerate(data,1):
        buckcell(t.cell(i,0),f"B{i}",i-1,size=BF)
        cell(t.cell(i,1),[(cn,{'sz':BF})],align=C,mgn=0.02)
        cell(t.cell(i,2),[(pc,{'sz':BF})],align=C,mgn=0.02)
        cell(t.cell(i,3),[(pj,{'sz':BF})],align=C,mgn=0.02)
        cell(t.cell(i,4),[(sc,{'sz':BF})],align=C,mgn=0.02)
    cell(t.cell(8,0),[("Tot",{'sz':BF})],bold=True,fill=TOT_BG,color=TOT_INK,mgn=0.02)
    for j,v in enumerate(tot,1): cell(t.cell(8,j),[(v,{'sz':BF,'b':True})],align=C,fill=TOT_BG,color=TOT_INK,mgn=0.02)
    t.rows[0].height=Inches(0.28)
    for i in range(1,9): t.rows[i].height=Inches(0.233)


BF=9.5
LX,LW,RX,RW=0.46,6.13,6.74,6.13
TY,TH,BY=1.24,2.74,4.06

def slide_resumen(meses, scn, n_estr, tasa, tasa_tot, v2, v2tot, v3, v3tot, tipos, v1, v1tot):
    s=add_slide()
    header(s,[(f"Buckets · {meses} ",{'c':GREEN_DEEP,'b':True}),("| Tasa de riesgo y distribución V1 / V2 / V3",{'c':GRAY_TITLE,'b':True})],
           [(f"Escenario {scn} · 7 buckets sobre {n_estr} estrategias del Árbol 2 · verde = menor riesgo, rojo = mayor",{'c':GRAY_TXT})], scn)
    # --- Tasa (arriba-izquierda) ---
    ix,iy,iw=card(s,LX,TY,LW,TH,"Tasa de riesgo por bucket")
    t=s.shapes.add_table(9,6,Inches(ix),Inches(iy),Inches(iw),Inches(2.144)).table
    t.first_row=False; t.horz_banding=False
    ws=[0.5, 1.42, 0.95, 0.95, 1.05, iw-0.5-1.42-0.95-0.95-1.05]
    setw(t,ws)
    th_row(t,0,[("Buck.",C,''),("Estrat.",L,''),("Tasa",C,''),("Prob.",C,''),("N",C,''),("% tot",C,'')])
    for i,(est,ta,pj,n,pt) in enumerate(tasa,1):
        buckcell(t.cell(i,0),f"B{i}",i-1,size=BF)
        cell(t.cell(i,1),[(est,{'sz':BF})],align=L,mgn=0.02)
        cell(t.cell(i,2),[(ta,{'sz':BF})],align=C,mgn=0.02)
        cell(t.cell(i,3),[(pj,{'sz':BF})],align=C,mgn=0.02)
        cell(t.cell(i,4),[(n,{'sz':BF})],align=C,mgn=0.02)
        cell(t.cell(i,5),[(pt,{'sz':BF})],align=C,mgn=0.02)
    cell(t.cell(8,0),[("Tot",{'sz':BF})],bold=True,fill=TOT_BG,color=TOT_INK,mgn=0.02)
    cell(t.cell(8,1),[(tasa_tot[0],{'sz':BF,'b':True})],align=L,fill=TOT_BG,color=TOT_INK,mgn=0.02)
    for j,v in enumerate(tasa_tot[1:],2): cell(t.cell(8,j),[(v,{'sz':BF,'b':True})],align=C,fill=TOT_BG,color=TOT_INK,mgn=0.02)
    t.rows[0].height=Inches(0.28)
    for i in range(1,9): t.rows[i].height=Inches(0.233)
    # --- V2 (arriba-derecha) · V3 (abajo-izquierda) ---
    dist_table(s,RX,TY,RW,"· Base inicial",v2,v2tot)
    dist_table(s,LX,BY,LW,"· Fuera de campaña",v3,v3tot)
    # --- V1 pivote (abajo-derecha) ---
    ncol=1+len(tipos)+1
    ix,iy,iw=card(s,RX,BY,RW,TH,"· Base de campañas",hcolor=GREEN_TEAL,sub="· periodo 20260624 · tipo × bucket")
    tv=s.shapes.add_table(9,ncol,Inches(ix),Inches(iy),Inches(iw),Inches(2.18)).table
    tv.first_row=False; tv.horz_banding=False
    bw=0.44; totw=0.92; tw=(iw-bw-totw)/len(tipos)
    setw(tv,[bw]+[tw]*len(tipos)+[totw])
    cell(tv.cell(0,0),[("Buck.",{'sz':10})],color=WHITE,bold=True,fill=TH_BG,align=C)
    for j,tp in enumerate(tipos,1): cell(tv.cell(0,j),[(tp,{'sz':8.3})],color=WHITE,bold=True,fill=TH_BG,align=C)
    cell(tv.cell(0,ncol-1),[("Total",{'sz':10})],color=TOTC_INK,bold=True,fill=TOTC_BG,align=C)
    for i,row in enumerate(v1,1):
        buckcell(tv.cell(i,0),f"B{i}",i-1,size=BF)
        for j in range(len(tipos)): cell(tv.cell(i,1+j),[(row[j],{'sz':BF})],align=C,mgn=0.02)
        cell(tv.cell(i,ncol-1),[(row[len(tipos)],{'sz':BF,'b':True})],align=C,fill=TOTC_BG,color=TOTC_INK,mgn=0.02)
    cell(tv.cell(8,0),[("Tot",{'sz':BF})],bold=True,fill=TOT_BG,color=TOT_INK,mgn=0.02)
    for j in range(len(tipos)): cell(tv.cell(8,1+j),[(v1tot[j],{'sz':BF,'b':True})],align=C,fill=TOT_BG,color=TOT_INK,mgn=0.02)
    cell(tv.cell(8,ncol-1),[(v1tot[-1],{'sz':BF,'b':True})],align=C,fill=TOT_BG,color=TOT_INK,mgn=0.02)
    tv.rows[0].height=Inches(0.38)
    for i in range(1,9): tv.rows[i].height=Inches(0.225)
    footer(s,[("Tasa malos:"," % de target_60_12m=1 en la base de modelamiento (RD, n=45,459)."),
              ("V1:"," base de campañas por tipo · V2: base de generación · V3: fuera de campaña."),
              ("Prob. / score:"," promedios por bucket · Rechazo: sin bucket asignado.")])


def slide_reglas(meses, scn, n_estr, reglas, rowh, varsline):
    s=add_slide()
    header(s,[(f"Buckets · {meses} ",{'c':GREEN_DEEP,'b':True}),("| Reglas por bucket",{'c':GRAY_TITLE,'b':True})],
           [(f"Escenario {scn} · {n_estr} estrategias del Árbol 2 agrupadas en 7 buckets · cada regla = una hoja del árbol",{'c':GRAY_TXT})], scn)
    nrows=1+sum(len(g[2]) for g in reglas)
    TYr=1.28; hh0=0.26
    ch=0.41+hh0+rowh*(nrows-1)+0.14
    ix,iy,iw=card(s,0.46,TYr,12.41,ch,"Reglas por bucket",sub="· estrategia → condición de la hoja → tasa de malos")
    t=s.shapes.add_table(nrows,6,Inches(ix),Inches(iy),Inches(iw),Inches(hh0+rowh*(nrows-1))).table; t.first_row=False; t.horz_banding=False
    ws=[0.5,0.4,iw-0.5-0.4-0.8-0.65-0.85,0.8,0.65,0.85]
    setw(t,ws)
    th_row(t,0,[("Buck.",C,''),("#",C,''),("Regla (condición de la hoja)",L,''),("Tasa",C,''),("% tot",C,''),("N",C,'')])
    r=1
    for (b,hue,rules) in reglas:
        r0=r
        for (idx,rule,ta,pc,n) in rules:
            cell(t.cell(r,1),[(idx,{'b':True,'sz':9,'c':RGBColor(0x0A,0x5A,0x44)})],mgn=0.02)
            cell(t.cell(r,2),[(rule,{'f':MONO,'sz':8.8,'c':MONO_INK})],align=L,nowrap=True,mgn=0.02)
            cell(t.cell(r,3),[(ta,{'sz':9})],align=C,mgn=0.02); cell(t.cell(r,4),[(pc,{'sz':9})],align=C,mgn=0.02); cell(t.cell(r,5),[(n,{'sz':9})],align=C,mgn=0.02)
            r+=1
        a=t.cell(r0,0)
        if r-1>r0: a.merge(t.cell(r-1,0))
        buckcell(a,b,hue,size=10,mgn=0.02)
    t.rows[0].height=Inches(hh0)
    for i in range(1,nrows): t.rows[i].height=Inches(rowh)
    footer(s,[("Regla:"," condición de la hoja del Árbol 2 (variables binarizadas con OptBinning)."),
              ("Tasa:"," % de malos de la estrategia · % tot / N: peso y tamaño en la base."),
              ("Variables:"," "+varsline)])


slide_resumen("24 meses","24m_sinedad_6","21",TASA_24,TASA_24_TOT,V2_24,V2_24_TOT,V3_24,V3_24_TOT,V1_24_TIPOS,V1_24,V1_24_TOT)
slide_reglas("24 meses","24m_sinedad_6","21",REGLAS_24,0.228,"segmentacion_gdp_v2, rk_ing_num, meses_desde_primer_castigo.")
slide_resumen("5 años","5años_sinedad_6","15",TASA_5,TASA_5_TOT,V2_5,V2_5_TOT,V3_5,V3_5_TOT,V1_5_TIPOS,V1_5,V1_5_TOT)
slide_reglas("5 años","5años_sinedad_6","15",REGLAS_5,0.30,"segmentacion_gdp_v2, rk_ing_num, nro_entidades_castigo, deuda_cas.")

prs.save("/home/user/Controling/presentacion_buckets_v1v2v3.pptx")
print("OK pptx buckets guardado")
