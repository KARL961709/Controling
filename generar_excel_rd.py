# -*- coding: utf-8 -*-
"""
Genera un Excel con una grafica de lineas: una curva por categoria de bin.
Eje X = periodos (codmes_ejec), Eje Y = tasa de malos / RD (tasa de eventos).
Colores graduales: rojo = bin de mayor riesgo -> verde = bin de menor riesgo.
"""
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import Series
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.marker import Marker

# Datos crudos: (periodo, bin, tasa_malos)
datos = [
    (202309, "[682.50, inf)", 0.04178273), (202309, "[627.50, 682.50)", 0.101648352),
    (202309, "[570.50, 627.50)", 0.133595285), (202309, "[523.50, 570.50)", 0.185990338),
    (202309, "[426.50, 523.50)", 0.275316456), (202309, "(-inf, 426.50)", 0.637651822),
    (202310, "[682.50, inf)", 0.06006006), (202310, "[627.50, 682.50)", 0.092178771),
    (202310, "[570.50, 627.50)", 0.16008316), (202310, "[523.50, 570.50)", 0.171875),
    (202310, "[426.50, 523.50)", 0.25), (202310, "(-inf, 426.50)", 0.533653846),
    (202311, "[682.50, inf)", 0.06), (202311, "[570.50, 627.50)", 0.08908686),
    (202311, "[627.50, 682.50)", 0.101092896), (202311, "[523.50, 570.50)", 0.194029851),
    (202311, "[426.50, 523.50)", 0.308447937), (202311, "(-inf, 426.50)", 0.559888579),
    (202312, "[682.50, inf)", 0.054151625), (202312, "[627.50, 682.50)", 0.1),
    (202312, "[570.50, 627.50)", 0.135359116), (202312, "[523.50, 570.50)", 0.18694362),
    (202312, "[426.50, 523.50)", 0.242366412), (202312, "(-inf, 426.50)", 0.560117302),
    (202401, "[682.50, inf)", 0.058441558), (202401, "[627.50, 682.50)", 0.092741935),
    (202401, "[570.50, 627.50)", 0.119113573), (202401, "[523.50, 570.50)", 0.190114068),
    (202401, "[426.50, 523.50)", 0.209476309), (202401, "(-inf, 426.50)", 0.51026393),
    (202402, "[682.50, inf)", 0.074433657), (202402, "[627.50, 682.50)", 0.125),
    (202402, "[570.50, 627.50)", 0.134292566), (202402, "[523.50, 570.50)", 0.18),
    (202402, "[426.50, 523.50)", 0.261569416), (202402, "(-inf, 426.50)", 0.598820059),
    (202403, "[682.50, inf)", 0.041420118), (202403, "[570.50, 627.50)", 0.107969152),
    (202403, "[627.50, 682.50)", 0.116352201), (202403, "[523.50, 570.50)", 0.2),
    (202403, "[426.50, 523.50)", 0.259183673), (202403, "(-inf, 426.50)", 0.551020408),
    (202404, "[682.50, inf)", 0.04), (202404, "[570.50, 627.50)", 0.118257261),
    (202404, "[627.50, 682.50)", 0.121447028), (202404, "[523.50, 570.50)", 0.131578947),
    (202404, "[426.50, 523.50)", 0.225806452), (202404, "(-inf, 426.50)", 0.518817204),
    (202405, "[682.50, inf)", 0.077306733), (202405, "[627.50, 682.50)", 0.096491228),
    (202405, "[570.50, 627.50)", 0.132352941), (202405, "[523.50, 570.50)", 0.19760479),
    (202405, "[426.50, 523.50)", 0.262327416), (202405, "(-inf, 426.50)", 0.510928962),
    (202406, "[682.50, inf)", 0.066489362), (202406, "[627.50, 682.50)", 0.094512195),
    (202406, "[570.50, 627.50)", 0.120454545), (202406, "[523.50, 570.50)", 0.160583942),
    (202406, "[426.50, 523.50)", 0.254789272), (202406, "(-inf, 426.50)", 0.4825),
    (202407, "[682.50, inf)", 0.081145585), (202407, "[627.50, 682.50)", 0.101580135),
    (202407, "[570.50, 627.50)", 0.132173913), (202407, "[523.50, 570.50)", 0.166030534),
    (202407, "[426.50, 523.50)", 0.25498008), (202407, "(-inf, 426.50)", 0.530201342),
    (202408, "[682.50, inf)", 0.036414566), (202408, "[627.50, 682.50)", 0.082524272),
    (202408, "[570.50, 627.50)", 0.165957447), (202408, "[523.50, 570.50)", 0.167070218),
    (202408, "[426.50, 523.50)", 0.188172043), (202408, "(-inf, 426.50)", 0.496828753),
    (202409, "[682.50, inf)", 0.062645012), (202409, "[627.50, 682.50)", 0.111363636),
    (202409, "[570.50, 627.50)", 0.143389199), (202409, "[523.50, 570.50)", 0.174825175),
    (202409, "[426.50, 523.50)", 0.228714524), (202409, "(-inf, 426.50)", 0.610766046),
    (202410, "[682.50, inf)", 0.054545455), (202410, "[627.50, 682.50)", 0.089058524),
    (202410, "[570.50, 627.50)", 0.146236559), (202410, "[523.50, 570.50)", 0.154639175),
    (202410, "[426.50, 523.50)", 0.229205176), (202410, "(-inf, 426.50)", 0.550761421),
    (202411, "[682.50, inf)", 0.051094891), (202411, "[627.50, 682.50)", 0.109958506),
    (202411, "[523.50, 570.50)", 0.157258065), (202411, "[570.50, 627.50)", 0.158862876),
    (202411, "[426.50, 523.50)", 0.210616438), (202411, "(-inf, 426.50)", 0.662162162),
    (202412, "[682.50, inf)", 0.087613293), (202412, "[627.50, 682.50)", 0.122596154),
    (202412, "[570.50, 627.50)", 0.128070175), (202412, "[523.50, 570.50)", 0.187782805),
    (202412, "[426.50, 523.50)", 0.271719039), (202412, "(-inf, 426.50)", 0.620347395),
    (202501, "[682.50, inf)", 0.082644628), (202501, "[627.50, 682.50)", 0.103825137),
    (202501, "[570.50, 627.50)", 0.128440367), (202501, "[523.50, 570.50)", 0.183727034),
    (202501, "[426.50, 523.50)", 0.191056911), (202501, "(-inf, 426.50)", 0.554455446),
    (202502, "[682.50, inf)", 0.07967033), (202502, "[627.50, 682.50)", 0.099730458),
    (202502, "[570.50, 627.50)", 0.134065934), (202502, "[523.50, 570.50)", 0.180758017),
    (202502, "[426.50, 523.50)", 0.217687075), (202502, "(-inf, 426.50)", 0.540880503),
]

# Orden de bins de MENOR a MAYOR riesgo (para leyenda ordenada)
# menor riesgo -> mayor riesgo
bins_orden = [
    "[682.50, inf)",
    "[627.50, 682.50)",
    "[570.50, 627.50)",
    "[523.50, 570.50)",
    "[426.50, 523.50)",
    "(-inf, 426.50)",
]

# Colores graduales: verde (menor riesgo) -> rojo (mayor riesgo)
colores = {
    "[682.50, inf)":     "00B050",  # verde
    "[627.50, 682.50)":  "92D050",  # verde claro
    "[570.50, 627.50)":  "FFC000",  # amarillo/ambar
    "[523.50, 570.50)":  "FF9900",  # naranja
    "[426.50, 523.50)":  "FF5050",  # rojo claro
    "(-inf, 426.50)":    "C00000",  # rojo
}

# Periodos ordenados
periodos = sorted({d[0] for d in datos})

# Pivot: dict[bin][periodo] = tasa
tabla = {b: {} for b in bins_orden}
for periodo, b, tasa in datos:
    tabla[b][periodo] = tasa

# ---- Construccion del Excel ----
wb = Workbook()
ws = wb.active
ws.title = "RD por bin"

# Cabecera
ws.cell(row=1, column=1, value="codmes_ejec")
for j, b in enumerate(bins_orden, start=2):
    ws.cell(row=1, column=j, value=b)

# Filas de datos (una fila por periodo)
for i, periodo in enumerate(periodos, start=2):
    ws.cell(row=i, column=1, value=periodo)
    for j, b in enumerate(bins_orden, start=2):
        val = tabla[b].get(periodo)
        c = ws.cell(row=i, column=j, value=val)
        if val is not None:
            c.number_format = "0.0%"

n_filas = len(periodos)
last_row = 1 + n_filas

# ---- Grafica ----
chart = LineChart()
chart.title = "Ranking de Desempeno (RD) - Tasa de eventos por bin"
chart.style = 2
chart.y_axis.title = "Tasa de malos / RD"
chart.x_axis.title = "Periodo (codmes_ejec)"
chart.y_axis.numFmt = "0.0%"
chart.x_axis.delete = False
chart.y_axis.delete = False
chart.height = 12
chart.width = 26

# Categorias (eje X = periodos)
cats = Reference(ws, min_col=1, min_row=2, max_row=last_row)

# Datos: todas las columnas de bins (incluye cabecera como titulo)
data_ref = Reference(ws, min_col=2, max_col=1 + len(bins_orden),
                     min_row=1, max_row=last_row)
chart.add_data(data_ref, titles_from_data=True)
chart.set_categories(cats)

# Estilo por serie (orden de series == orden de bins_orden)
for serie, b in zip(chart.series, bins_orden):
    color = colores[b]
    serie.graphicalProperties.line = LineProperties(solidFill=color, w=28000)
    serie.smooth = False
    serie.marker = Marker(symbol="circle", size=5)
    serie.marker.graphicalProperties.solidFill = color
    serie.marker.graphicalProperties.line.solidFill = color

ws.add_chart(chart, "I2")

salida = "RD_por_bin.xlsx"
wb.save(salida)
print("Guardado:", salida)
