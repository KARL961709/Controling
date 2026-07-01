# -*- coding: utf-8 -*-
"""PPTX nativo y editable del deck Árbol 2 (6 slides): por escenario resumen,
reglas y CUBO de reglas (formas nativas: rects + triángulos freeform + textos)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

GREEN_DEEP=RGBColor(0x1F,0x51,0x30); GREEN_DEEP2=RGBColor(0x26,0x60,0x3A)
GREEN=RGBColor(0x1A,0xA4,0x4C); GREEN_BRIGHT=RGBColor(0x1A,0x98,0x50)
GREEN_TEAL=RGBColor(0x3A,0x7D,0x5C); GREEN_SOFT=RGBColor(0xE8,0xF5,0xEE)
TEAL=RGBColor(0x00,0x7A,0x72)
TH_BG=RGBColor(0x0A,0x5A,0x44); GRAY_TITLE=RGBColor(0x9A,0xA6,0xA0)
GRAY_TXT=RGBColor(0x5B,0x6F,0x66); INK=RGBColor(0x14,0x27,0x1D)
MONO_INK=RGBColor(0x23,0x43,0x3C); WHITE=RGBColor(0xFF,0xFF,0xFF)
CARD_LINE=RGBColor(0xE7,0xEE,0xE9); BORDER="DDE7E4"
TOT_BG=RGBColor(0xCF,0xE9,0xE4); TOT_INK=RGBColor(0x08,0x4B,0x44)
RJ_BG=RGBColor(0xEC,0xEF,0xF0); RJ_INK=RGBColor(0x4A,0x5A,0x56)
TOTC_BG=RGBColor(0xDF,0xF0,0xED); TOTC_INK=RGBColor(0x0A,0x5A,0x52)
STATBG=RGBColor(0xF5,0xFA,0xF7); STATLINE=RGBColor(0xE3,0xEF,0xE8)
CELLLINE=RGBColor(0xAE,0xBF,0xBA); FACE1=RGBColor(0xE8,0xEF,0xEC); FACE2=RGBColor(0xDD,0xE6,0xE2)
FONT="Segoe UI"; MONO="Consolas"
EMU=914400

def hsl(h,s=0.60,l=0.78):
    c=(1-abs(2*l-1))*s; x=c*(1-abs((h/60.0)%2-1)); m=l-c/2
    if   h<60:  r,g,b=c,x,0
    elif h<120: r,g,b=x,c,0
    elif h<180: r,g,b=0,c,x
    elif h<240: r,g,b=0,x,c
    elif h<300: r,g,b=x,0,c
    else:       r,g,b=c,0,x
    return RGBColor(int((r+m)*255),int((g+m)*255),int((b+m)*255))
HUE=[130,118,106,95,83,71,59,47,35,24,12,0]
BUCK=[hsl(h) for h in HUE]

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]
def add_slide(): return prs.slides.add_slide(BLANK)

def rrect(s,x,y,w,h,fill,line=None,radius=0.12,lw=1.0):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    try: sh.adjustments[0]=max(0.0,min(0.5,radius/min(w,h)))
    except Exception: pass
    sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    sh.shadow.inherit=False; return sh
def rect(s,x,y,w,h,fill,line=None,lw=0.75):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    sh.shadow.inherit=False; return sh
def poly(s,pts,fill,line=None,lw=0.75):
    P=[(int(round(x*EMU)),int(round(y*EMU))) for x,y in pts]   # EMU enteros, scale=1 (evita redondeo del path)
    fb=s.shapes.build_freeform(P[0][0],P[0][1],scale=1)
    fb.add_line_segments(P[1:],close=True)
    sh=fb.convert_to_shape()
    sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    sh.shadow.inherit=False; return sh
def tb(s,x,y,w,h,runs,size=12,color=INK,bold=False,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,font=FONT,ls=1.0,rot=0):
    bx=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));
    if rot: bx.rotation=rot
    tf=bx.text_frame; tf.word_wrap=True
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0;tf.vertical_anchor=anchor
    p=tf.paragraphs[0]; p.alignment=align
    if ls:p.line_spacing=ls
    if isinstance(runs,str):runs=[(runs,{})]
    for t,st in runs:
        r=p.add_run();r.text=t;r.font.size=Pt(st.get('sz',size));r.font.bold=st.get('b',bold)
        r.font.name=st.get('f',font);r.font.color.rgb=st.get('c',color)
    return bx
def set_border(cell,color=BORDER,w=9525):
    tcPr=cell._tc.get_or_add_tcPr()
    for tag in ('a:lnB','a:lnT','a:lnR','a:lnL'):
        for el in tcPr.findall(qn(tag)): tcPr.remove(el)
        ln=tcPr.makeelement(qn(tag),{'w':str(w),'cap':'flat'})
        sf=tcPr.makeelement(qn('a:solidFill'),{}); cl=tcPr.makeelement(qn('a:srgbClr'),{'val':color})
        sf.append(cl);ln.append(sf);tcPr.insert(0,ln)
def cell(c,runs,size=11,color=INK,bold=False,align=PP_ALIGN.CENTER,fill=None,font=FONT,anchor=MSO_ANCHOR.MIDDLE,mgn=0.025,nowrap=False,border=True):
    if fill is not None: c.fill.solid(); c.fill.fore_color.rgb=fill
    else: c.fill.background()
    c.vertical_anchor=anchor; c.margin_left=Inches(0.05); c.margin_right=Inches(0.04)
    c.margin_top=Inches(mgn); c.margin_bottom=Inches(mgn)
    tf=c.text_frame; tf.word_wrap=(not nowrap)
    p=tf.paragraphs[0]; p.alignment=align
    if isinstance(runs,str):runs=[(runs,{})]
    first=True
    for t,st in runs:
        r=p.runs[0] if(first and p.runs) else p.add_run(); first=False
        r.text=t;r.font.size=Pt(st.get('sz',size));r.font.bold=st.get('b',bold)
        r.font.name=st.get('f',font);r.font.color.rgb=st.get('c',color)
    if border: set_border(c)
def card(s,x,y,w,h,title,hcolor=GREEN_DEEP,sub=None,hh=0.34):
    rrect(s,x,y,w,h,WHITE,line=CARD_LINE,radius=0.12)
    rrect(s,x,y,w,hh,hcolor,radius=0.12); rect(s,x,y+hh-0.13,w,0.13,hcolor)
    runs=[(title,{'b':True,'sz':13,'c':WHITE})]
    if sub: runs.append(("   "+sub,{'sz':10.5,'c':WHITE}))
    tb(s,x+0.15,y,w-0.3,hh,runs,anchor=MSO_ANCHOR.MIDDLE)
    return x+0.14,y+hh+0.07,w-0.28
def header(s,title_runs,subtitle,badge):
    tb(s,0.46,0.24,10.0,0.5,title_runs,size=24,bold=True,anchor=MSO_ANCHOR.MIDDLE)
    tb(s,0.46,0.76,10.4,0.3,subtitle,size=11,color=GRAY_TXT,anchor=MSO_ANCHOR.MIDDLE)
    bw=2.5; rrect(s,13.333-0.46-bw,0.26,bw,0.42,GREEN_SOFT,radius=0.21)
    tb(s,13.333-0.46-bw,0.26,bw,0.42,[("●  ",{'c':GREEN,'b':True,'sz':10}),(badge,{'c':GREEN_DEEP,'b':True,'sz':11})],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
def footer(s,items):
    rect(s,0.46,7.02,13.333-0.92,0.012,STATLINE)
    runs=[]
    for i,(lead,rest) in enumerate(items):
        if i:runs.append(("      ",{}))
        runs.append((lead,{'b':True,'c':GREEN_DEEP2,'sz':8})); runs.append((rest,{'c':GRAY_TITLE,'sz':8}))
    tb(s,0.46,7.08,13.333-0.92,0.34,runs,size=8,color=GRAY_TITLE)
def th_row(t,row,headers,size=9.5):
    for j,(htxt,al,sp) in enumerate(headers):
        fill=TH_BG;col=WHITE
        if sp=='rj':fill=RJ_BG;col=RJ_INK
        if sp=='tot':fill=TOTC_BG;col=TOTC_INK
        cell(t.cell(row,j),htxt,size=size,color=col,bold=True,fill=fill,align=al,nowrap=True)
def buckcell(c,label,hidx,size=10,mgn=0.02):
    cell(c,label,size=size,color=INK,bold=True,fill=BUCK[hidx],align=PP_ALIGN.CENTER,mgn=mgn)

def P(d):   # decimal -> % (1 decimal, .0 sin decimal)
    v=round(float(d)*1000)/10
    return (f"{v:.0f}" if abs(v-round(v))<1e-9 else f"{v:.1f}")+"%"
def S(x): return str(round(float(x)))                      # score entero
def R(s):   # % existente, redondea .0
    v=round(float(str(s).replace('%',''))*10)/10
    return (f"{v:.0f}" if abs(v-round(v))<1e-9 else f"{v:.1f}")+"%"

# ===================== DATOS =====================
exec(open("/home/user/Controling/_cubo_data.py").read())

# ===================== SLIDE RESUMEN =====================
def slide_resumen(meses,scn,nestr,tasa,tasaT,v2,v2T,v3,v3T,tipos,tipoT):
    s=add_slide()
    header(s,[(f"Buckets Árbol 2 · {meses} ",{'c':GREEN_DEEP,'b':True}),("| Tasa de riesgo y distribución V1 / V2 / V3",{'c':GRAY_TITLE,'b':True})],
           [(f"Escenario {scn} · 12 buckets sobre {nestr} estrategias del Árbol 2 (3 variables) · verde = menor riesgo, rojo = mayor",{'c':GRAY_TXT})],scn)
    C=PP_ALIGN.CENTER; L=PP_ALIGN.LEFT
    LX,LW,RX,RW=0.46,6.13,6.74,6.13
    TY,TH,BY=1.12,2.88,4.10
    FS=7; HS=7.5; rh=0.165
    # --- Tasa (arriba-izq) ---
    ix,iy,iw=card(s,LX,TY,LW,TH,"Tasa de riesgo por bucket")
    t=s.shapes.add_table(14,6,Inches(ix),Inches(iy),Inches(iw),Inches(0.19+rh*13)).table; t.first_row=False;t.horz_banding=False
    ws=[0.42,1.4,0.85,0.85,1.05,iw-0.42-1.4-0.85-0.85-1.05]
    for j,wd in enumerate(ws):t.columns[j].width=Inches(wd)
    th_row(t,0,[("Buck.",C,''),("Estrat.",L,''),("Tasa",C,''),("Prob.",C,''),("N",C,''),("% tot",C,'')],size=HS)
    for i,r in enumerate(tasa,1):
        buckcell(t.cell(i,0),r[0],i-1,size=FS,mgn=0.012); cell(t.cell(i,1),[(r[1],{'sz':FS})],align=L,nowrap=True,mgn=0.012)
        cell(t.cell(i,2),[(r[2],{'sz':FS})],align=C,nowrap=True,mgn=0.012);cell(t.cell(i,3),[(P(r[5]),{'sz':FS})],align=C,nowrap=True,mgn=0.012)
        cell(t.cell(i,4),[(r[3],{'sz':FS})],align=C,nowrap=True,mgn=0.012);cell(t.cell(i,5),[(R(r[4]),{'sz':FS})],align=C,nowrap=True,mgn=0.012)
    cell(t.cell(13,0),"Tot",size=HS,bold=True,fill=TOT_BG,color=TOT_INK,nowrap=True,mgn=0.012)
    tvals=[(tasaT[0],L),(tasaT[1],C),(P(tasaT[4]),C),(tasaT[2],C),(R(tasaT[3]),C)]
    for j,(v,al) in enumerate(tvals,1):cell(t.cell(13,j),[(v,{'b':True,'sz':FS})],align=al,fill=TOT_BG,color=TOT_INK,nowrap=True,mgn=0.012)
    t.rows[0].height=Inches(0.19)
    for i in range(1,14):t.rows[i].height=Inches(rh)
    # --- V2 (arriba-der) · V3 (abajo-izq) ---
    for (cx,cy,ttl,data,tot) in [(RX,TY,"· Base inicial",v2,v2T),(LX,BY,"· Fuera de campaña",v3,v3T)]:
        ix,iy,iw=card(s,cx,cy,LW,TH,ttl,hcolor=GREEN_BRIGHT)
        tt=s.shapes.add_table(14,5,Inches(ix),Inches(iy),Inches(iw),Inches(0.19+rh*13)).table; tt.first_row=False;tt.horz_banding=False
        ws2=[0.42,1.45,0.95,1.15,iw-0.42-1.45-0.95-1.15]
        for j,wd in enumerate(ws2):tt.columns[j].width=Inches(wd)
        th_row(tt,0,[("Buck.",C,''),("Cantidad",C,''),("%",C,''),("Prob.",C,''),("score",C,'')],size=HS)
        for i,r in enumerate(data,1):
            buckcell(tt.cell(i,0),f"B{i}",i-1,size=FS,mgn=0.012)
            for j,val in enumerate([r[0],R(r[1]),P(r[2]),S(r[3])]):cell(tt.cell(i,j+1),[(val,{'sz':FS})],align=C,nowrap=True,mgn=0.012)
        cell(tt.cell(13,0),"Tot",size=HS,bold=True,fill=TOT_BG,color=TOT_INK,nowrap=True,mgn=0.012)
        for j,val in enumerate([tot[0],R(tot[1]),P(tot[2]),S(tot[3])]):cell(tt.cell(13,j+1),[(val,{'b':True,'sz':FS})],align=C,fill=TOT_BG,color=TOT_INK,nowrap=True,mgn=0.012)
        tt.rows[0].height=Inches(0.19)
        for i in range(1,14):tt.rows[i].height=Inches(rh)
    # --- V1 pivote por BUCKET (abajo-der): Buck | tipo1..N | Total ---
    nt=len(tipos); ncol=1+nt+1; VF=6.5
    ix,iy,iw=card(s,RX,BY,RW,TH,"· Base de campañas",hcolor=GREEN_TEAL,sub="· 20260624 · tipo × bucket")
    tv=s.shapes.add_table(14,ncol,Inches(ix),Inches(iy),Inches(iw),Inches(0.34+0.155*13)).table; tv.first_row=False;tv.horz_banding=False
    bw=0.38; totw=0.85; tw=(iw-bw-totw)/nt
    for j,wd in enumerate([bw]+[tw]*nt+[totw]):tv.columns[j].width=Inches(wd)
    cell(tv.cell(0,0),[("Buck.",{'sz':HS})],color=WHITE,bold=True,fill=TH_BG,align=C)
    for j,tp in enumerate(tipos,1):cell(tv.cell(0,j),[(tp[0],{'sz':6})],color=WHITE,bold=True,fill=TH_BG,align=C)
    cell(tv.cell(0,ncol-1),[("Total",{'sz':HS})],color=TOTC_INK,bold=True,fill=TOTC_BG,align=C)
    for k in range(1,13):
        buckcell(tv.cell(k,0),f"B{k}",k-1,size=VF,mgn=0.01)
        for j in range(nt):cell(tv.cell(k,1+j),[(tipos[j][k],{'sz':VF})],align=C,nowrap=True,mgn=0.01)
        cell(tv.cell(k,ncol-1),[(tipoT[k-1],{'sz':VF,'b':True})],align=C,fill=TOTC_BG,color=TOTC_INK,nowrap=True,mgn=0.01)
    cell(tv.cell(13,0),[("Tot",{'sz':VF})],bold=True,fill=TOT_BG,color=TOT_INK,nowrap=True,mgn=0.01)
    for j in range(nt):cell(tv.cell(13,1+j),[(tipos[j][13],{'sz':VF,'b':True})],align=C,fill=TOT_BG,color=TOT_INK,nowrap=True,mgn=0.01)
    cell(tv.cell(13,ncol-1),[(tipoT[12],{'sz':VF,'b':True})],align=C,fill=TOT_BG,color=TOT_INK,nowrap=True,mgn=0.01)
    tv.rows[0].height=Inches(0.34)
    for k in range(1,14):tv.rows[k].height=Inches(0.155)
    footer(s,[("Tasa malos:"," % de target_60_12m=1 en la base de modelamiento (RD, n=45,459)."),
              ("V1:"," base de campañas por tipo (20260624) · V2: base de generación · V3: fuera de campaña."),
              ("Prob. / score:"," promedios por bucket.")])

# ===================== SLIDE REGLAS =====================
def reglas_tbl(s,x,y,w,rows,off,rowh=0.225):
    n=len(rows)+1
    t=s.shapes.add_table(n,6,Inches(x),Inches(y),Inches(w),Inches(0.26+rowh*(n-1))).table; t.first_row=False;t.horz_banding=False
    ws=[0.34,0.30,w-0.34-0.30-0.62-0.5-0.62,0.62,0.5,0.62]
    for j,wd in enumerate(ws):t.columns[j].width=Inches(wd)
    th_row(t,0,[("Bk",PP_ALIGN.CENTER,''),("#",PP_ALIGN.CENTER,''),("Regla",PP_ALIGN.LEFT,''),("Tasa",PP_ALIGN.CENTER,''),("%",PP_ALIGN.CENTER,''),("N",PP_ALIGN.CENTER,'')])
    for i,r in enumerate(rows,1):
        buckcell(t.cell(i,0),f"B{r[0]}",r[0]-1,size=8,mgn=0.015)
        cell(t.cell(i,1),[(str(off+i),{'b':True,'sz':8.5,'c':RGBColor(0x0A,0x5A,0x44)})],mgn=0.015)
        cell(t.cell(i,2),[(r[1],{'f':MONO,'sz':8.3,'c':MONO_INK})],align=PP_ALIGN.LEFT,nowrap=True,mgn=0.015)
        cell(t.cell(i,3),[(R(r[2]),{'sz':8.5})],mgn=0.015);cell(t.cell(i,4),[(R(r[3]),{'sz':8.5})],align=PP_ALIGN.RIGHT,mgn=0.015);cell(t.cell(i,5),[(r[4],{'sz':8.5})],align=PP_ALIGN.RIGHT,mgn=0.015)
    t.rows[0].height=Inches(0.24)
    for i in range(1,n):t.rows[i].height=Inches(rowh)
    return n

def slide_reglas(meses,scn,nreglas,rows,dos):
    s=add_slide()
    header(s,[(f"Buckets Árbol 2 · {meses} ",{'c':GREEN_DEEP,'b':True}),("| Reglas por bucket",{'c':GRAY_TITLE,'b':True})],
           [(f"Escenario {scn} · {nreglas} estrategias del Árbol 2 · cada regla = una hoja",{'c':GRAY_TXT})],scn)
    if dos:
        ix,iy,iw=card(s,0.46,1.30,12.41,5.45,"Reglas (estrategia → condición → tasa de malos)")
        half=(iw-0.3)/2
        reglas_tbl(s,ix,iy,half,rows[:15],0,rowh=0.225)
        reglas_tbl(s,ix+half+0.3,iy,half,rows[15:],15,rowh=0.225)
    else:
        n=len(rows)
        ix,iy,iw=card(s,0.46,1.30,12.41,0.55+0.30*n,"Reglas (estrategia → condición → tasa de malos)")
        reglas_tbl(s,ix,iy,iw,rows,0,rowh=0.29)
    footer(s,[("pcc"," = puntuacion_cal_cat (Score Rebank) · saldo = saldo_prom_tot_pasivo_u3m (Saldo Pasivo) · bk = bucket (Árbol 1)."),
              ("Bk:"," bucket final · %: % del total · N: clientes.")])

# ===================== SLIDE CUBO =====================
def slide_cubo(meses,scn,nreglas,grid,rows,cols,umbral,desc,tasa,scats,ccats):
    s=add_slide()
    header(s,[(f"Buckets Árbol 2 · {meses} ",{'c':GREEN_DEEP,'b':True}),("| Cubo de reglas (3 variables)",{'c':GRAY_TITLE,'b':True})],
           [(f"Las {nreglas} reglas en 3 variables: ",{'c':GRAY_TXT}),("Score Rebank × bucket Árbol 1 × Saldo Pasivo",{'c':GRAY_TXT,'b':True}),(" → color = bucket final",{'c':GRAY_TXT})],scn)
    ix,iy,iw=card(s,0.46,1.34,12.41,5.18,"Mapa de reglas en el cubo",sub="· cara = Score Rebank × bucket Árbol 1 · profundidad = Saldo Pasivo")
    # --- cubo ---
    nr=len(rows);nc=len(cols);cs=0.52;dd=0.24;dyy=-0.17
    ox=2.42;oy=2.42;GW=nc*cs;GH=nr*cs
    poly(s,[(ox,oy),(ox+GW,oy),(ox+GW+dd,oy+dyy),(ox+dd,oy+dyy)],FACE1,line=RGBColor(0xCD,0xD9,0xD4))
    poly(s,[(ox+GW,oy),(ox+GW,oy+GH),(ox+GW+dd,oy+GH+dyy),(ox+GW+dd,oy+dyy)],FACE2,line=RGBColor(0xCD,0xD9,0xD4))
    for r in range(nr):
        for c in range(nc):
            x=ox+c*cs;y=oy+r*cs;g=grid[r][c]
            if isinstance(g,tuple):
                poly(s,[(x,y),(x+cs,y),(x+cs,y+cs)],BUCK[g[0]-1],line=WHITE,lw=1)
                poly(s,[(x,y),(x,y+cs),(x+cs,y+cs)],BUCK[g[1]-1],line=WHITE,lw=1)
                tb(s,x+cs*0.40,y+cs*0.04,cs*0.58,cs*0.34,f"B{g[0]}",size=8.5,bold=True,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
                tb(s,x+cs*0.02,y+cs*0.60,cs*0.58,cs*0.34,f"B{g[1]}",size=8.5,bold=True,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
            else:
                rect(s,x,y,cs,cs,BUCK[g-1],line=WHITE,lw=1)
                tb(s,x,y,cs,cs,f"B{g}",size=11,bold=True,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
            rect(s,x,y,cs,cs,None,line=CELLLINE,lw=0.5)
    # etiquetas filas (Score Rebank): categoría + umbral (corresponde con la tabla)
    for r in range(nr):
        tb(s,ox-1.32,oy+r*cs+0.04,1.20,0.18,[(scats[r],{'sz':8,'b':True,'c':MONO_INK})],align=PP_ALIGN.RIGHT,anchor=MSO_ANCHOR.MIDDLE)
        tb(s,ox-1.32,oy+r*cs+0.25,1.20,0.16,[(rows[r],{'sz':7,'c':GRAY_TXT,'f':MONO})],align=PP_ALIGN.RIGHT,anchor=MSO_ANCHOR.MIDDLE)
    # etiquetas columnas (bucket) + zonas de segmento
    for c in range(nc):tb(s,ox+c*cs,oy+GH+0.02,cs,0.18,cols[c],size=10,color=MONO_INK,bold=True,align=PP_ALIGN.CENTER)
    c0=0
    while c0<nc:
        e=c0
        while e+1<nc and ccats[e+1]==ccats[c0]: e+=1
        tb(s,ox+c0*cs,oy+GH+0.20,(e-c0+1)*cs,0.16,[("seg. "+ccats[c0],{'sz':8,'b':True,'c':GRAY_TXT})],align=PP_ALIGN.CENTER)
        c0=e+1
    # eje vertical (nombre rotado)
    vcx=0.82;vcy=oy+GH/2
    tb(s,vcx-0.70,vcy-0.12,1.40,0.24,"Score Rebank",size=11.5,color=TEAL,bold=True,align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE,rot=270)
    # eje horizontal = segmento
    tb(s,ox,oy+GH+0.42,GW,0.22,"bucket Árbol 1 · segmento",size=11.5,color=TEAL,bold=True,align=PP_ALIGN.CENTER)
    tb(s,ox,oy+GH+0.63,GW,0.2,"segmento bajo → alto  (menor → mayor riesgo)",size=9,color=GRAY_TXT,align=PP_ALIGN.CENTER)
    # profundidad (Saldo Pasivo) con categorías + umbral
    tb(s,ox+GW*0.24,oy+dyy-0.36,GW*0.7+0.9,0.18,"Saldo Pasivo",size=11,color=TEAL,bold=True,align=PP_ALIGN.CENTER)
    tb(s,ox+GW*0.24,oy+dyy-0.19,GW*0.7+0.9,0.16,f"▲ alto (> {umbral})  ·  ▼ bajo (≤ {umbral})",size=8,color=GRAY_TXT,align=PP_ALIGN.CENTER)
    # --- panel derecho ---
    px=7.05;pw=5.6
    # mini-clave saldo
    rrect(s,px,iy,pw,0.78,STATBG,line=STATLINE,radius=0.08)
    ex,ey=px+0.13,iy+0.13
    poly(s,[(ex,ey),(ex+0.42,ey),(ex+0.42,ey+0.42)],hsl(106),line=WHITE,lw=1)
    poly(s,[(ex,ey),(ex,ey+0.42),(ex+0.42,ey+0.42)],hsl(83),line=WHITE,lw=1)
    rect(s,ex,ey,0.42,0.42,None,line=CELLLINE,lw=0.5)
    tb(s,px+0.66,iy+0.08,pw-0.78,0.62,
       [("La profundidad del cubo es el Saldo Pasivo. Si cambia el resultado, la casilla se parte en diagonal:  ",{'c':INK,'sz':9.5}),
        ("▲ Saldo Pasivo alto",{'c':INK,'b':True,'sz':9.5}),(f" (> {umbral})  ·  ",{'c':GRAY_TXT,'sz':9.5}),
        ("▼ Saldo Pasivo bajo",{'c':INK,'b':True,'sz':9.5}),(f" (≤ {umbral})",{'c':GRAY_TXT,'sz':9.5})],
       size=9.5,color=INK,anchor=MSO_ANCHOR.MIDDLE,ls=1.05)
    tb(s,px,iy+0.86,pw,0.22,"Bucket final → ¿qué cliente es?",size=12.5,color=GREEN_DEEP,bold=True)
    lt=s.shapes.add_table(12,3,Inches(px),Inches(iy+1.10),Inches(pw),Inches(0.255*12)).table; lt.first_row=False;lt.horz_banding=False
    lt.columns[0].width=Inches(0.42);lt.columns[1].width=Inches(0.62);lt.columns[2].width=Inches(pw-0.42-0.62)
    for i in range(12):
        buckcell(lt.cell(i,0),f"B{i+1}",i,size=9.5,mgn=0.02)
        cell(lt.cell(i,1),[(tasa[i][2],{'sz':9.5})],align=PP_ALIGN.RIGHT,color=GRAY_TXT)
        cell(lt.cell(i,2),[(desc[i],{'sz':9.3,'c':MONO_INK})],align=PP_ALIGN.LEFT)
    for i in range(12):lt.rows[i].height=Inches(0.255)
    tb(s,px,iy+1.10+0.255*12+0.06,pw,0.4,
       [("Orden por riesgo: ",{'c':GREEN_DEEP,'b':True,'sz':10}),("domina el Score Rebank; dentro de cada tramo de score sube el segmento del Árbol 1.",{'c':GRAY_TXT,'sz':10})],
       size=10,color=GRAY_TXT,ls=1.1)
    footer(s,[("Score Rebank:"," puntuacion_cal_cat (1–6) · bucket Árbol 1: bucket (ordinal) · Saldo Pasivo: saldo_prom_tot_pasivo_u3m."),
              ("Umbral Saldo Pasivo:"," "+umbral+".")])

# ===================== MONTАЖ =====================
slide_resumen("24 meses","24m_sinedad_6","30",TASA24,TASA24T,V224,V224T,V324,V324T,V1P24,V1P24T)
slide_reglas("24 meses","24m_sinedad_6","30",REG24,True)
slide_cubo("24 meses","24m_sinedad_6","30",CUBE24,ROWS24,COLS24,"349.07",DESC24,TASA24,SCATS24,CCATS24)
slide_resumen("5 años","5años_sinedad_4","16",TASA5,TASA5T,V25,V25T,V35,V35T,V1P5,V1P5T)
slide_reglas("5 años","5años_sinedad_4","16",REG5,False)
slide_cubo("5 años","5años_sinedad_4","16",CUBE5,ROWS5,COLS5,"220.89",DESC5,TASA5,SCATS5,CCATS5)

prs.save("/home/user/Controling/presentacion_buckets_arbol2_cubo.pptx")
print("OK pptx cubo guardado")
