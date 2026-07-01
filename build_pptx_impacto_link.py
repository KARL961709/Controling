# -*- coding: utf-8 -*-
"""PPT de 1 slide: imagen del simulador (snapshot a pantalla completa) +
botón con hipervínculo que abre el HTML interactivo. La imagen también es clicable."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

LINK = "impacto_leads_slide.html"   # ruta relativa: mantener el HTML junto al PPTX
IMG  = "impacto_leads_snapshot.png"
GREEN=RGBColor(0x1A,0xA4,0x4C); GREEN_DEEP=RGBColor(0x0A,0x5A,0x34); WHITE=RGBColor(0xFF,0xFF,0xFF)

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
s=prs.slides.add_slide(prs.slide_layouts[6])

# imagen a pantalla completa
pic=s.shapes.add_picture(IMG,Inches(0),Inches(0),Inches(13.333),Inches(7.5))
# toda la imagen es un hipervínculo
pic.click_action.hyperlink.address=LINK

# botón visible con hipervínculo (banda blanca libre entre la tabla y el pie)
bw,bh=3.40,0.50; bx=(13.333-bw)/2; by=6.46
btn=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(bx),Inches(by),Inches(bw),Inches(bh))
btn.adjustments[0]=0.5
btn.fill.solid(); btn.fill.fore_color.rgb=GREEN
btn.line.color.rgb=GREEN_DEEP; btn.line.width=Pt(1.25); btn.shadow.inherit=False
tf=btn.text_frame; tf.word_wrap=False; tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
tf.vertical_anchor=MSO_ANCHOR.MIDDLE
p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER
r=p.add_run(); r.text="▶  Abrir simulador interactivo (HTML)"; r.font.size=Pt(14); r.font.bold=True
r.font.name="Segoe UI"; r.font.color.rgb=WHITE
btn.click_action.hyperlink.address=LINK

prs.save("presentacion_impacto_link.pptx")
print("OK presentacion_impacto_link.pptx  (link ->", LINK, ")")
