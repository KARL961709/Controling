# -*- coding: utf-8 -*-
"""PPTX nativo y editable (2 slides):
   1) Variables de entrada por árbol y escenario (lenguaje de negocio)
   2) Propuesta de seguimiento del proceso (ciclo + matriz + cadencia)
Formas nativas: rects + ovals + tabla + textos. Nada de imágenes."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

GREEN_DEEP=RGBColor(0x1F,0x51,0x30); GREEN_DEEP2=RGBColor(0x26,0x60,0x3A)
GREEN=RGBColor(0x1A,0xA4,0x4C); GREEN_BRIGHT=RGBColor(0x1A,0x98,0x50)
GREEN_TEAL=RGBColor(0x2F,0x7D,0x70); GREEN_SOFT=RGBColor(0xE8,0xF5,0xEE)
GRAY_TITLE=RGBColor(0x9A,0xA6,0xA0); GRAY_TXT=RGBColor(0x5B,0x6F,0x66)
INK=RGBColor(0x14,0x27,0x1D); WHITE=RGBColor(0xFF,0xFF,0xFF)
CARD_LINE=RGBColor(0xE7,0xEE,0xE9); SOFT=RGBColor(0xF6,0xFA,0xF8)
AMBER=RGBColor(0xE7,0x9A,0x1E); RED=RGBColor(0xD2,0x4B,0x3E)
AMBER_BG=RGBColor(0xFD,0xF1,0xDC); RED_BG=RGBColor(0xFB,0xE4,0xE1)
CHIP_BG=RGBColor(0xE9,0xEE,0xFB); CHIP_INK=RGBColor(0x34,0x52,0x9C); CHIP_LINE=RGBColor(0xD4,0xDD,0xF6)
FOOT_LINE=RGBColor(0xE3,0xEC,0xE7); FOOT_INK=RGBColor(0x8A,0x9B,0x93)
ROWALT=RGBColor(0xF7,0xFB,0xF9); BORDER="EAF1ED"
# tags de familia
T_SEG=(RGBColor(0xE5,0xF0,0xFF),RGBColor(0x2B,0x5A,0xA6))
T_ING=(RGBColor(0xEA,0xFB,0xE9),RGBColor(0x2F,0x8A,0x34))
T_CAS=(RGBColor(0xFF,0xE9,0xE3),RGBColor(0xB5,0x50,0x2F))
T_SAL=(RGBColor(0xFF,0xF3,0xD9),RGBColor(0x9A,0x6A,0x12))
T_SCO=(RGBColor(0xEF,0xE7,0xFF),RGBColor(0x6A,0x45,0xA0))
T_BUC=(RGBColor(0xE2,0xF4,0xEF),RGBColor(0x0A,0x7A,0x6F))
FONT="Segoe UI"; EMU=914400

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]

def rrect(s,x,y,w,h,fill,line=None,radius=0.12,lw=1.0):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    try: sh.adjustments[0]=max(0.0,min(0.5,radius/min(w,h)))
    except Exception: pass
    sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    sh.shadow.inherit=False; return sh
def rect(s,x,y,w,h,fill,line=None,lw=0.75):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    if fill is None: sh.fill.background()
    else: sh.fill.solid(); sh.fill.fore_color.rgb=fill
    if line is None: sh.line.fill.background()
    else: sh.line.color.rgb=line; sh.line.width=Pt(lw)
    sh.shadow.inherit=False; return sh
def oval(s,x,y,w,h,fill):
    sh=s.shapes.add_shape(MSO_SHAPE.OVAL,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=fill; sh.line.fill.background(); sh.shadow.inherit=False; return sh
def hline(s,x1,x2,y,color,w):
    ln=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y),Inches(x2),Inches(y))
    ln.line.color.rgb=color; ln.line.width=Pt(w); ln.shadow.inherit=False; return ln
def tb(s,x,y,w,h,runs,size=12,color=INK,bold=False,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,ls=1.05):
    bx=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=bx.text_frame; tf.word_wrap=True
    tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0;tf.vertical_anchor=anchor
    p=tf.paragraphs[0]; p.alignment=align
    if ls:p.line_spacing=ls
    if isinstance(runs,str):runs=[(runs,{})]
    for t,st in runs:
        r=p.add_run();r.text=t;r.font.size=Pt(st.get('sz',size));r.font.bold=st.get('b',bold)
        r.font.name=FONT;r.font.color.rgb=st.get('c',color)
    return bx
def chip(s,x,y,txt,size=7.5,bg=CHIP_BG,ink=CHIP_INK,line=CHIP_LINE,padx=0.055):
    w=0.10+len(txt)*(size*0.0092)+padx
    rrect(s,x,y,w,0.20,bg,line=line,radius=0.09,lw=0.75)
    tb(s,x,y,w,0.20,[(txt,{'sz':size,'b':True,'c':ink})],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    return w

def header(s,title_main,title_sec,sub_runs,badge):
    tb(s,0.42,0.30,10.6,0.42,[(title_main+" ",{'sz':19,'b':True,'c':GREEN_DEEP}),
                              (title_sec,{'sz':19,'b':True,'c':GRAY_TITLE})],anchor=MSO_ANCHOR.MIDDLE)
    tb(s,0.42,0.76,10.8,0.28,sub_runs)
    bw=0.28+len(badge)*0.078
    rrect(s,12.91-bw,0.34,bw,0.34,GREEN_SOFT,radius=0.17)
    oval(s,12.91-bw+0.14,0.45,0.11,0.11,GREEN)
    tb(s,12.91-bw+0.30,0.34,bw-0.3,0.34,[(badge,{'sz':9,'b':True,'c':GREEN_DEEP})],anchor=MSO_ANCHOR.MIDDLE)

def footer(s,items):
    hline(s,0.42,12.91,7.02,FOOT_LINE,1.0)
    x=0.42
    for lead,rest,w in items:
        tb(s,x,7.10,w,0.34,[(lead+" ",{'sz':7.6,'b':True,'c':GREEN_DEEP2}),(rest,{'sz':7.6,'c':FOOT_INK})])
        x+=w+0.25

# ============================================================ SLIDE 1
s=prs.slides.add_slide(BLANK)
header(s,"Variables de entrada por árbol y escenario","| Score Rebank",
       [("Qué variables alimentan cada árbol, descritas en lenguaje de negocio · el ",{'sz':10,'c':GRAY_TXT}),
        ("Árbol 1",{'sz':10,'b':True,'c':GREEN_DEEP}),(" segmenta y su resultado entra al ",{'sz':10,'c':GRAY_TXT}),
        ("Árbol 2",{'sz':10,'b':True,'c':GREEN_DEEP})],"24 meses y 5 años")

def varbox(x,y,w,num,title,role,color,cols,note_runs,res_chips):
    boxH=2.86
    rrect(s,x,y,w,boxH,WHITE,line=CARD_LINE,radius=0.11,lw=1.0)
    # header
    hh=0.44
    rrect(s,x,y,w,hh,color,radius=0.11); rect(s,x,y+hh-0.11,w,0.11,color)
    rrect(s,x+0.12,y+0.11,0.24,0.24,RGBColor(0x5F,0x9E,0x83),radius=0.06)
    tb(s,x+0.12,y+0.11,0.24,0.24,[(str(num),{'sz':10,'b':True,'c':WHITE})],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    tb(s,x+0.44,y,w-2.6,hh,[(title,{'sz':13.5,'b':True,'c':WHITE})],anchor=MSO_ANCHOR.MIDDLE)
    tb(s,x+w-2.5,y,2.38,hh,[(role,{'sz':9.5,'b':True,'c':WHITE})],align=PP_ALIGN.RIGHT,anchor=MSO_ANCHOR.MIDDLE)
    # dos columnas
    colw=(w)/2
    bodyTop=y+hh+0.10
    for ci,(pill,cnt,items) in enumerate(cols):
        cx=x+ci*colw+0.14
        if ci==1: rect(s,x+colw,y+hh,0.008,boxH-hh-0.52,RGBColor(0xEE,0xF3,0xF0))
        wp=chip(s,cx,bodyTop,pill,size=8.5,bg=GREEN_DEEP2,ink=WHITE,line=None)
        tb(s,cx+wp+0.08,bodyTop,1.4,0.20,[(cnt,{'sz':9.5,'b':True,'c':GREEN_DEEP})],anchor=MSO_ANCHOR.MIDDLE)
        yy=bodyTop+0.32
        for tag,(tbg,tink),runs in items:
            wt=0.12+len(tag)*0.058
            rrect(s,cx,yy+0.012,wt,0.18,tbg,radius=0.05)
            tb(s,cx,yy+0.012,wt,0.18,[(tag,{'sz':7.5,'b':True,'c':tink})],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
            tb(s,cx+wt+0.07,yy,colw-wt-0.34,0.40,runs,size=9.3,ls=1.06)
            yy+=0.40
    # nota inferior
    ny=y+boxH-0.44
    rect(s,x,ny,w,0.44,SOFT); rect(s,x,ny,w,0.012,RGBColor(0xEE,0xF3,0xF0))
    tb(s,x+0.14,ny,w*0.52,0.44,note_runs,size=9.5,anchor=MSO_ANCHOR.MIDDLE)
    rx=x+w*0.55
    for k,v in res_chips:
        txt=k+"  "+v
        wc=0.18+len(txt)*0.052
        rrect(s,rx,ny+0.10,wc,0.24,WHITE,line=RGBColor(0xE2,0xEC,0xE7),radius=0.07,lw=0.75)
        tb(s,rx+0.09,ny+0.10,wc-0.12,0.24,[(k+"  ",{'sz':8,'c':GRAY_TXT}),(v,{'sz':8,'b':True,'c':GREEN_DEEP})],anchor=MSO_ANCHOR.MIDDLE)
        rx+=wc+0.10

BW=6.12; L=0.42; R=0.42+BW+0.25; TY=1.14
varbox(L,TY,BW,1,"Árbol 1 · segmentación base","variables del cliente",GREEN_TEAL,
 [("24 meses","3 variables",[
    ("segmento",T_SEG,[("Segmento comercial del cliente",{'b':True,'c':GREEN_DEEP}),(" (grupos G1 a G8).",{})]),
    ("ingreso",T_ING,[("Ranking de ingreso",{'b':True,'c':GREEN_DEEP}),(" del cliente.",{})]),
    ("castigo",T_CAS,[("Meses desde el primer castigo",{'b':True,'c':GREEN_DEEP}),(".",{})]),
  ]),
  ("5 años","4 variables",[
    ("segmento",T_SEG,[("Segmento comercial del cliente",{'b':True,'c':GREEN_DEEP}),(" (grupos G1 a G8).",{})]),
    ("ingreso",T_ING,[("Ranking de ingreso",{'b':True,'c':GREEN_DEEP}),(" del cliente.",{})]),
    ("castigo",T_CAS,[("Nº de entidades con castigo",{'b':True,'c':GREEN_DEEP}),(".",{})]),
    ("castigo",T_CAS,[("Monto de deuda castigada",{'b':True,'c':GREEN_DEEP}),(".",{})]),
  ])],
 [("No usa saldo pasivo",{'sz':9.5,'b':True,'c':GREEN_DEEP}),(" · solo perfil y comportamiento de castigo.",{'sz':9.5,'c':GRAY_TXT})],
 [("resultado","7 buckets → Árbol 2")])

varbox(R,TY,BW,2,"Árbol 2 · refinamiento","bucket + saldo + score",GREEN_BRIGHT,
 [("24 meses","3 variables",[
    ("bucket",T_BUC,[("Bucket del Árbol 1",{'b':True,'c':GREEN_DEEP}),(" (segmento base ordenado por riesgo).",{})]),
    ("saldo",T_SAL,[("Saldo pasivo promedio",{'b':True,'c':GREEN_DEEP}),(" de los últimos 3 meses.",{})]),
    ("score",T_SCO,[("Categoría de Score Rebank",{'b':True,'c':GREEN_DEEP}),(" (de 1 a 6).",{})]),
  ]),
  ("5 años","3 variables",[
    ("bucket",T_BUC,[("Bucket del Árbol 1",{'b':True,'c':GREEN_DEEP}),(" (segmento base ordenado por riesgo).",{})]),
    ("saldo",T_SAL,[("Saldo pasivo promedio",{'b':True,'c':GREEN_DEEP}),(" de los últimos 3 meses.",{})]),
    ("score",T_SCO,[("Categoría de Score Rebank",{'b':True,'c':GREEN_DEEP}),(" (de 1 a 6).",{})]),
  ])],
 [("Mismas variables",{'sz':9.5,'b':True,'c':GREEN_DEEP}),(" en ambos escenarios · aquí sí entra el saldo.",{'sz':9.5,'c':GRAY_TXT})],
 [("24m","30 → 12"),("5a","16 → 12")])

# adicionales
def addcard(x,w,ic,title,bullets,chips=None):
    y=4.24; h=1.86
    rrect(s,x,y,w,h,WHITE,line=CARD_LINE,radius=0.11,lw=1.0)
    rrect(s,x,y,w,0.36,GREEN_DEEP,radius=0.11); rect(s,x,y+0.25,w,0.11,GREEN_DEEP)
    rrect(s,x+0.11,y+0.09,0.19,0.19,RGBColor(0x4A,0x7A,0x62),radius=0.05)
    tb(s,x+0.11,y+0.09,0.19,0.19,[(ic,{'sz':10,'b':True,'c':WHITE})],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    tb(s,x+0.36,y,w-0.4,0.36,[(title,{'sz':11,'b':True,'c':WHITE})],anchor=MSO_ANCHOR.MIDDLE)
    yy=y+0.46
    for runs in bullets:
        oval(s,x+0.13,yy+0.05,0.05,0.05,GREEN)
        tb(s,x+0.26,yy,w-0.40,0.5,runs,size=9,ls=1.08)
        yy+=0.335
    if chips:
        cx=x+0.24
        for c in chips:
            cx+=chip(s,cx,yy,c,size=7.5)+0.06

AW=(12.49-2*0.14)/3
addcard(L,AW,"↕","Cómo entra cada variable",[
  [("Cada variable entra con su ",{}),("dirección de riesgo esperada",{'b':True,'c':GREEN_DEEP}),(".",{})],
  [("Los valores se ",{}),("agrupan en tramos ordenados por riesgo",{'b':True,'c':GREEN_DEEP}),(".",{})],
  [("Cliente ",{}),("sin dato",{'b':True,'c':GREEN_DEEP}),(" → tramo de mayor riesgo.",{})]])
addcard(L+AW+0.14,AW,"✓","Qué se conserva",[
  [("Solo las ",{}),("variables que aportan",{'b':True,'c':GREEN_DEEP}),(" (el resto se descarta).",{})],
  [("Una sola variable de saldo",{'b':True,'c':GREEN_DEEP}),(" — solo en el Árbol 2.",{})],
  [("Métodos de ",{}),("ordenamiento",{'b':True,'c':GREEN_DEEP}),(" probados:",{})]],
  chips=["Cascada","ChiMerge","OptBin","Shrinkage","MDLP"])
addcard(L+2*(AW+0.14),AW,"◎","Variables evaluadas (Árbol 1)",[
  [("Comportamiento de castigo:",{'b':True,'c':GREEN_DEEP}),(" deuda, montos, entidades, meses.",{})],
  [("Ingreso y perfil:",{'b':True,'c':GREEN_DEEP}),(" ranking, sexo, nivel profesional, segmento.",{})],
  [("La ",{}),("edad",{'b':True,'c':GREEN_DEEP}),(" se prueba en una versión aparte.",{})]])

footer(s,[("Árbol 1 → Árbol 2:","el bucket del primer árbol es una entrada del segundo · el saldo pasivo solo aparece en el Árbol 2.",6.05),
          ("Escenarios:","24 meses y 5 años cambian variables y nº de estrategias, no el procedimiento.",5.9)])

# ============================================================ SLIDE 2
s=prs.slides.add_slide(BLANK)
header(s,"Propuesta de seguimiento del proceso","| Score Rebank",
       [("Cómo vigilar en el tiempo que los ",{'sz':10,'c':GRAY_TXT}),("árboles",{'sz':10,'b':True,'c':GREEN_DEEP}),
        (" y las ",{'sz':10,'c':GRAY_TXT}),("campañas de rebancarización",{'sz':10,'b':True,'c':GREEN_DEEP}),
        (" siguen funcionando, y cuándo actuar",{'sz':10,'c':GRAY_TXT})],"monitoreo continuo")

# ciclo de vida
cyc=[("1 · Construcción","Árbol 1 → Árbol 2 → buckets finales",GREEN_DEEP),
     ("2 · Despliegue","buckets aplicados a campañas / pilotos",GREEN_TEAL),
     ("3 · Seguimiento","estabilidad · poder · negocio",GREEN_BRIGHT),
     ("4 · Decisión","mantener · recalibrar · reentrenar",GREEN_DEEP2)]
cy0=1.18; cyh=0.56; nw=2.86; gap=0.34; cx=0.42
for i,(t,d,col) in enumerate(cyc):
    rrect(s,cx,cy0,nw,cyh,col,radius=0.10)
    tb(s,cx+0.12,cy0+0.06,nw-0.2,0.24,[(t,{'sz':11.5,'b':True,'c':WHITE})])
    tb(s,cx+0.12,cy0+0.30,nw-0.2,0.22,[(d,{'sz':8.7,'c':WHITE})])
    if i<3:
        tb(s,cx+nw,cy0,gap,cyh,[("→",{'sz':16,'b':True,'c':RGBColor(0x8F,0xB3,0xA3)})],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    cx+=nw+gap
tb(s,0.42,cy0+cyh+0.02,12.49,0.22,[("Ciclo cerrado: ",{'sz':9,'b':True,'c':GREEN_DEEP}),
   ("la Decisión realimenta la Construcción — el seguimiento es lo que dispara recalibrar o reentrenar.",{'sz':9,'c':GRAY_TXT})],align=PP_ALIGN.CENTER)

# ---- matriz (tabla nativa) ----
def set_border(c,color=BORDER,w=9525):
    tcPr=c._tc.get_or_add_tcPr()
    for tag in ('a:lnB','a:lnT','a:lnR','a:lnL'):
        for el in tcPr.findall(qn(tag)): tcPr.remove(el)
        ln=tcPr.makeelement(qn(tag),{'w':str(w),'cap':'flat'})
        sf=tcPr.makeelement(qn('a:solidFill'),{}); cl=tcPr.makeelement(qn('a:srgbClr'),{'val':color})
        sf.append(cl);ln.append(sf);tcPr.insert(0,ln)
def setcell(c,paras,fill=None,anchor=MSO_ANCHOR.MIDDLE):
    if fill is not None: c.fill.solid(); c.fill.fore_color.rgb=fill
    else: c.fill.background()
    c.vertical_anchor=anchor; c.margin_left=Inches(0.07); c.margin_right=Inches(0.06)
    c.margin_top=Inches(0.03); c.margin_bottom=Inches(0.03)
    tf=c.text_frame; tf.word_wrap=True
    for i,runs in enumerate(paras):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.line_spacing=1.02; p.space_after=Pt(1)
        for t,st in runs:
            r=p.add_run(); r.text=t; r.font.size=Pt(st.get('sz',8)); r.font.bold=st.get('b',False)
            r.font.name=FONT; r.font.color.rgb=st.get('c',INK)
    set_border(c)

HEAD=["Dimensión","Qué vigila","Métrica","Umbral de alerta","Acción"]
ROWS=[
 ("Estabilidad de la población","Que la mezcla de clientes por bucket no se desvíe de la base de construcción.",
  "PSI de la distribución de buckets",[("ámbar > 0.10",AMBER,AMBER_BG),("rojo > 0.25",RED,RED_BG)],"Revisar la mezcla y recalibrar los cortes."),
 ("Estabilidad de las variables","Que cada variable de entrada mantenga su distribución en el tiempo.",
  "PSI por variable",[("ámbar > 0.10",AMBER,AMBER_BG),("rojo > 0.25",RED,RED_BG)],"Re-agrupar los tramos de la variable."),
 ("Orden y poder del modelo","Que los buckets sigan ordenados por riesgo y separando bien.",
  "Monotonía de la tasa de malos · KS / Gini",[("se rompe el orden",AMBER,AMBER_BG),("caída de KS / Gini",RED,RED_BG)],"Reentrenar el árbol."),
 ("Calibración del riesgo","Que la tasa de malos observada coincida con la esperada por bucket.",
  "Observado vs. esperado",[("desvío moderado",AMBER,AMBER_BG),("desvío alto sostenido",RED,RED_BG)],"Recalibrar los puntos de corte."),
 ("Desempeño de negocio","Que la rebancarización genere valor frente a un grupo de control.",
  "Activación · mora temprana · recupero",[("bajo el objetivo",AMBER,AMBER_BG),("mora sobre lo tolerado",RED,RED_BG)],"Ajustar la estrategia comercial."),
 ("Cobertura y rechazo","Cuántos clientes objetivo capturan las reglas y cuántos quedan fuera.",
  "% capturado · % de rechazo",[("cae la cobertura",AMBER,AMBER_BG),("rechazo creciente",RED,RED_BG)],"Revisar o ampliar las reglas."),
]
TX=0.42; TW=12.49; TTY=2.34
gf=s.shapes.add_table(len(ROWS)+1,5,Inches(TX),Inches(TTY),Inches(TW),Inches(0.3))
tbl=gf.table; tbl.first_row=False; tbl.horz_banding=False
# desactivar estilo por defecto
tbl_pr=tbl._tbl.tblPr
tbl_pr.set('firstRow','0'); tbl_pr.set('bandRow','0')
colw=[0.17,0.27,0.20,0.16,0.20]
for i,frac in enumerate(colw): tbl.columns[i].width=Inches(TW*frac)
tbl.rows[0].height=Inches(0.30)
for r in range(1,len(ROWS)+1): tbl.rows[r].height=Inches(0.46)
for i,h in enumerate(HEAD):
    setcell(tbl.cell(0,i),[[(h,{'sz':9,'b':True,'c':WHITE})]],fill=GREEN_DEEP)
for ri,(dim,vig,met,thr,act) in enumerate(ROWS,1):
    fill=ROWALT if ri%2==0 else WHITE
    setcell(tbl.cell(ri,0),[[(dim,{'sz':9,'b':True,'c':GREEN_DEEP})]],fill=fill)
    setcell(tbl.cell(ri,1),[[(vig,{'sz':8.3,'c':INK})]],fill=fill)
    setcell(tbl.cell(ri,2),[[(met,{'sz':8.3,'b':True,'c':GREEN_DEEP2})]],fill=fill)
    setcell(tbl.cell(ri,3),[[(thr[0][0],{'sz':8,'b':True,'c':thr[0][1]})],
                            [(thr[1][0],{'sz':8,'b':True,'c':thr[1][1]})]],fill=fill)
    setcell(tbl.cell(ri,4),[[(act,{'sz':8.3,'b':True,'c':GREEN_DEEP2})]],fill=fill)

# ---- franja inferior ----
by=5.44; bh=1.48
# pane cadencia
pw1=7.55
rrect(s,0.42,by,pw1,bh,WHITE,line=CARD_LINE,radius=0.11,lw=1.0)
rrect(s,0.42,by,pw1,0.34,GREEN_DEEP,radius=0.11); rect(s,0.42,by+0.23,pw1,0.11,GREEN_DEEP)
tb(s,0.56,by,pw1-0.2,0.34,[("Cadencia propuesta",{'sz':11,'b':True,'c':WHITE})],anchor=MSO_ANCHOR.MIDDLE)
cad=[("Mensual",GREEN_TEAL,["Estabilidad (PSI población y variables)","Cobertura y rechazo","Desempeño de campañas"]),
     ("Trimestral",GREEN_BRIGHT,["Orden y poder (KS / Gini)","Calibración del riesgo","Monotonía por bucket"]),
     ("Anual / por alerta",GREEN_DEEP2,["Reentrenamiento completo de los árboles","Revisión del universo de variables"])]
cw=(pw1-0.28-2*0.10)/3
for i,(k,kc,items) in enumerate(cad):
    cxx=0.56+i*(cw+0.10); cyy=by+0.42
    rrect(s,cxx,cyy,cw,bh-0.52,RGBColor(0xF9,0xFC,0xFB),line=RGBColor(0xEE,0xF3,0xF0),radius=0.07,lw=0.75)
    wk=0.16+len(k)*0.066
    rrect(s,cxx+0.09,cyy+0.08,wk,0.20,kc,radius=0.09)
    tb(s,cxx+0.09,cyy+0.08,wk,0.20,[(k,{'sz':8.3,'b':True,'c':WHITE})],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    yy=cyy+0.34
    for it in items:
        oval(s,cxx+0.12,yy+0.045,0.045,0.045,GREEN)
        tb(s,cxx+0.24,yy,cw-0.32,0.30,[(it,{'sz':8,'c':INK})],ls=1.02)
        yy+=0.215
# pane semáforo/gobierno
px=0.42+pw1+0.20; pw2=12.91-px
rrect(s,px,by,pw2,bh,WHITE,line=CARD_LINE,radius=0.11,lw=1.0)
rrect(s,px,by,pw2,0.34,GREEN_DEEP,radius=0.11); rect(s,px,by+0.23,pw2,0.11,GREEN_DEEP)
tb(s,px+0.14,by,pw2-0.2,0.34,[("Semáforo y gobierno",{'sz':11,'b':True,'c':WHITE})],anchor=MSO_ANCHOR.MIDDLE)
sem=[("Verde","dentro de rango — sin acción.",RGBColor(0x2F,0xA8,0x63)),
     ("Ámbar","vigilar y documentar.",AMBER),
     ("Rojo","recalibrar o reentrenar.",RED)]
yy=by+0.42
for k,d,col in sem:
    oval(s,px+0.16,yy+0.02,0.13,0.13,col)
    tb(s,px+0.38,yy,pw2-0.5,0.22,[(k,{'sz':9,'b':True,'c':GREEN_DEEP}),("  · "+d,{'sz':9,'c':INK})],anchor=MSO_ANCHOR.MIDDLE)
    yy+=0.23
hline(s,px+0.16,px+pw2-0.16,yy+0.01,RGBColor(0xEE,0xF3,0xF0),1.0)
tb(s,px+0.16,yy+0.06,pw2-0.32,0.40,[("Roles: ",{'sz':8,'b':True,'c':GREEN_DEEP}),
   ("Riesgos define umbrales · Modelos ejecuta el monitoreo · Negocio actúa en campañas · toda decisión queda en bitácora.",{'sz':8,'c':GRAY_TXT})],ls=1.1)

footer(s,[("PSI:","mide cuánto se desplaza una distribución respecto a la base de construcción.",5.2),
          ("KS / Gini:","miden qué tan bien el modelo separa clientes buenos de malos.",4.2),
          ("Control:","grupo comparable que no recibe la campaña.",2.6)])

prs.save("presentacion_variables_seguimiento.pptx")
print("OK presentacion_variables_seguimiento.pptx")
