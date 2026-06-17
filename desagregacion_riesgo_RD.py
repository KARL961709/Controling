"""
====================================================================================
DESAGREGACIÓN DE RIESGO  (proxy = RD / tasa de default, continuo)
====================================================================================
Objetivo = RD (tasa de default). RD alto = MAYOR riesgo. Verde = RD bajo / rojo = RD alto.

Metodología A (Solo árbol): variables CRUDAS + restricción monótona de negocio.
Metodología B (Optbinning + árbol):
  - Se fija la dirección de negocio (+1 / -1) para que optbinning agrupe la variable
    en el sentido correcto.
  - El WoE de cada bin ENTRA al árbol (la variable se reemplaza por su WoE).
    -> como el WoE ya codifica el riesgo, la restricción en el árbol es +1 para todas.
  - En el árbol y en los cuadros se muestran los LABELS (rangos de los bins), NO el
    número de WoE.
  - Si una variable no logra binarizar (ni 5 ni 2 bins) -> se descarta.
  - Por variable se exporta la TABLA optbinning (bin / WoE / RD medio) en Excel y TXT.

El PUNTAJE entra como variable: banda G1..G5 (vía CORTES por situación laboral) -> 1..5.
Selección de leads por APETITO de RD (no top-%): un segmento entra si su RD <= apetito;
se evalúan escenarios apetito × {0.90, 0.95, 1.00, 1.05, 1.10}.

Requisitos: pip install pandas numpy scikit-learn matplotlib openpyxl scipy optbinning
====================================================================================
"""

import warnings
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from scipy.stats import spearmanr
from sklearn.tree import DecisionTreeRegressor, plot_tree
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

warnings.filterwarnings("ignore")

# ====================================================================================
# CONFIGURACIÓN  (edita esto)
# ====================================================================================
DATA_PATH = "dataprueba1606VF.csv"
RD_COL    = "rd"                            # *** TARGET: tasa de default (continuo) ***
SCORE_COL = "puntaje_mod"                   # puntaje -> entra como banda 1..5
SIT_COL   = "sit_laboral_mod"              # situación laboral (para bandear el puntaje)
FLG_COL   = "FLG_CAST_AP"
FAR_COL   = "flg_far_mto_trx_presencial_12m_c216"
VALORES_FLG = ["CAST_NOIBK_REP>=5anios", "NO_CAST_NOIBK_U24M"]

NUM_VARS = [
    "edad_num", "rk_ing_num", "DEUDA_CAS", "monto_castigado_total",
    "monto_castigado_otros", "nro_entidades_castigo", "max_dias_mora_castigo",
    "meses_desde_ultimo_castigo", "meses_desde_primer_castigo",
    "saldo_pasivo_actual", "saldo_prom_pasivo", "saldo_activo_actual",
    "prom_saldo_pasivo_u4m", "max_saldo_pasivo_u6m", "nro_meses_con_pasivo_u6m",
]
ORD_VAR = "segmentacion_gdp_v2"
SCORE_BANDA = "score_g"
VARS_EXCLUIR = ["edad_num", "rk_ing_num"]

WOE_AL_ARBOL = True                         # B: el WoE de los bins entra al árbol (labels en el árbol)
SENTINEL  = -99999999
MAX_DEPTH = 6
MIN_SAMPLES_LEAF = 0.03
OPTB_SAMPLE = 200_000
RANDOM_STATE = 42

CORTES = {
    "DEPENDIENTE":    [("G1", 967), ("G2", 941), ("G3", 876)],
    "MIXTO":          [("G1", 971), ("G2", 949), ("G3", 915), ("G4", 872), ("G5", 852)],
    "DEPEN_EXPERIAN": [("G1", 971), ("G2", 949), ("G3", 915), ("G4", 872), ("G5", 852)],
    "INDEPENDIENTE":  [("G1", 981), ("G2", 960), ("G3", 943), ("G4", 898), ("G5", 844)],
}

# Dirección monótona POR NEGOCIO frente al RD (RD alto = peor).
DIRECCION_NEGOCIO = {
    "edad_num": 0, "rk_ing_num": -1, "DEUDA_CAS": +1, "monto_castigado_total": +1,
    "monto_castigado_otros": +1, "nro_entidades_castigo": +1, "max_dias_mora_castigo": +1,
    "meses_desde_ultimo_castigo": -1, "meses_desde_primer_castigo": -1,
    "saldo_pasivo_actual": -1, "saldo_prom_pasivo": -1, "saldo_activo_actual": 0,
    "prom_saldo_pasivo_u4m": -1, "max_saldo_pasivo_u6m": -1, "nro_meses_con_pasivo_u6m": -1,
    "segmentacion_gdp_v2": +1, "score_g": +1,
}

SUBJECT_COL = "subject_id"
APETITO_RD = 0.02
FACTORES_RD = [0.90, 0.95, 1.00, 1.05, 1.10]
TOP_N_ESTRATEGIAS = 20
LEADS_EN_EXCEL_MAX = 100_000
BINS_OBJETIVO = [5, 2]

CMAP = plt.cm.RdYlGn
BORDER = Border(*[Side(style="thin", color="999999")] * 4)
BOLD = Font(bold=True)


# ====================================================================================
# UTILIDADES
# ====================================================================================
def norm_riesgo(v, vmin, vmax):
    """1 = verde = RD BAJO (menor riesgo)."""
    return 0.5 if vmax <= vmin else (vmax - v) / (vmax - vmin)


def hex_color(t):
    r, g, b, _ = CMAP(float(np.clip(t, 0, 1)))
    return f"{int(r * 255):02X}{int(g * 255):02X}{int(b * 255):02X}"


def banda_cortes_vec(score, sit):
    score = np.asarray(score, dtype=float)
    sit = np.asarray(sit, dtype=object)
    out = np.full(score.shape, "s/d", dtype=object)
    out[np.isnan(score)] = "Missing"
    for key, tabla in CORTES.items():
        m = (sit == key) & ~np.isnan(score)
        if not m.any():
            continue
        bandas = np.array([b for b, _ in tabla])
        umbrales = np.array([u for _, u in tabla], dtype=float)
        idx = (umbrales[None, :] > score[m][:, None]).sum(axis=1)
        out[m] = bandas[np.minimum(idx, len(tabla) - 1)]
    return out


def fmt_intervalo(lo, hi):
    if hi < -1e7:
        return "Missing"
    if lo < -1e7:
        lo = -np.inf
    if lo == -np.inf and hi == np.inf:
        return "(todos)"
    if lo == -np.inf:
        return f"<= {hi:,.0f}"
    if hi == np.inf:
        return f"> {lo:,.0f}"
    return f"({lo:,.0f}, {hi:,.0f}]"


def _parse_bin(s):
    """Convierte el texto de un bin de optbinning a (low, high) numéricos."""
    s = str(s).strip()
    if s in ("Special", "Missing", "", "nan"):
        return None
    s = s.strip("[]()")
    if "," not in s:
        return None
    a, b = s.split(",", 1)

    def f(z):
        z = z.strip()
        if "inf" in z:
            return -np.inf if z.startswith("-") else np.inf
        try:
            return float(z)
        except ValueError:
            return None
    lo, hi = f(a), f(b)
    return None if lo is None or hi is None else (lo, hi)


def labelmap_from_table(tb, valcol):
    """Construye el mapeo {valor_WoE_o_mean -> rango original} a partir de la tabla optbinning."""
    vbins, miss_val = [], None
    for _, row in tb.iterrows():
        b = str(row["Bin"])
        val = pd.to_numeric(row.get(valcol), errors="coerce")
        if b == "Missing":
            if pd.notna(val) and row.get("Count", 0) > 0:
                miss_val = float(val)
            continue
        pr = _parse_bin(b)
        if pr is None or pd.isna(val):
            continue
        vbins.append((float(val), pr[0], pr[1]))
    vbins.sort(key=lambda z: z[0])
    return {"vbins": vbins, "miss_val": miss_val}


def cond_label(woemap, var, lo, hi):
    """Etiqueta legible de la condición (lo, hi] de 'var'.
    - Variables WoE: traduce el rango de WoE al RANGO ORIGINAL de los bins (+ 'Missing' si aplica).
    - Variables crudas/ordinales: usa el intervalo real (fmt_intervalo)."""
    if var not in woemap:
        return fmt_intervalo(lo, hi)
    info = woemap[var]
    vb, mv = info["vbins"], info["miss_val"]
    sel = [(low, high) for (w, low, high) in vb if (w > lo and w <= hi)]
    inc_miss = (mv is not None and mv > lo and mv <= hi)
    total = len(vb) + (1 if mv is not None else 0)
    if len(sel) + (1 if inc_miss else 0) >= total:
        return "(todos)"
    partes = []
    if sel:
        partes.append(fmt_intervalo(min(l for l, _ in sel), max(h for _, h in sel)))
    if inc_miss:
        partes.append("Missing")
    return " + ".join(partes) if partes else "(ninguno)"


def reglas_por_hoja(tree, feats):
    t = tree.tree_
    out, stack = {}, [(0, {})]
    while stack:
        node, conds = stack.pop()
        if t.children_left[node] == -1:
            out[node] = conds
            continue
        f, thr = feats[t.feature[node]], t.threshold[node]
        lo, hi = conds.get(f, (-np.inf, np.inf))
        izq = dict(conds); izq[f] = (lo, min(hi, thr))
        der = dict(conds); der[f] = (max(lo, thr), hi)
        stack += [(t.children_left[node], izq), (t.children_right[node], der)]
    return out


def arbol_texto(tree, feats, woemap):
    """Árbol en texto mostrando LABELS (rangos), no umbrales WoE. value = RD del nodo."""
    t = tree.tree_
    lines = []

    def rec(node, depth):
        pre = "|   " * depth + "|--- "
        if t.children_left[node] == -1:
            lines.append("|   " * depth + f"|--- RD = {t.value[node][0][0]:.4f}")
            return
        f, thr = feats[t.feature[node]], t.threshold[node]
        lines.append(pre + f"{f}: {cond_label(woemap, f, -np.inf, thr)}")
        rec(t.children_left[node], depth + 1)
        lines.append(pre + f"{f}: {cond_label(woemap, f, thr, np.inf)}")
        rec(t.children_right[node], depth + 1)

    rec(0, 0)
    return "\n".join(lines)


def optbinning_fit(x, y, nombre, direccion):
    """Ajusta optbinning (target continuo) forzando dirección, exigiendo bins (5, luego 2).
    Devuelve (feasible, etiqueta, tabla, dir_efectiva, optb)."""
    from optbinning import ContinuousOptimalBinning
    trend = {1: "ascending", -1: "descending", 0: "auto_asc_desc"}[direccion]
    etiqueta = {1: "CRECIENTE RD (+)", -1: "DECRECIENTE RD (-)", 0: "auto"}[direccion]
    if (~np.isnan(x)).sum() < 50:
        return False, "sin datos -> descartada", None, 0, None
    optb = None
    for min_bins in BINS_OBJETIVO:
        try:
            o = ContinuousOptimalBinning(name=nombre, monotonic_trend=trend,
                                         min_n_bins=min_bins, max_n_bins=max(min_bins, 8))
            o.fit(x, y)
            if o.status in ("OPTIMAL", "FEASIBLE") and len(o.splits) >= (min_bins - 1):
                optb = o
                break
        except Exception:
            continue
    if optb is None:
        return False, etiqueta + " -> SIN BINARIZACIÓN (ni 5 ni 2 bins), descartada", None, 0, None
    n_bins = len(optb.splits) + 1
    tb = optb.binning_table.build()
    if direccion != 0:
        dir_ef = direccion
    else:
        mask = ~tb["Bin"].astype(str).isin(["Special", "Missing", ""])
        means = pd.to_numeric(tb.loc[mask, "Mean"], errors="coerce").dropna().values
        dir_ef = (1 if means[-1] >= means[0] else -1) if len(means) >= 2 else 0
        etiqueta = ("CRECIENTE RD (+)" if dir_ef > 0 else "DECRECIENTE RD (-)") + " (detectada)"
    return True, f"{etiqueta} [{n_bins} bins]", tb, dir_ef, optb


def escribir_cuadro(ws, r0, c0, titulo, tab, vfil, vcol, modo, vmin=None, vmax=None):
    ws.cell(r0, c0, titulo).font = Font(bold=True, size=12)
    ws.cell(r0 + 1, c0 + 1, vcol).font = BOLD
    ws.cell(r0 + 2, c0, vfil).font = BOLD
    for j, col in enumerate(tab.columns):
        c = ws.cell(r0 + 2, c0 + 1 + j, str(col))
        c.font = BOLD; c.border = BORDER
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    mx = np.nanmax(tab.values) if tab.size else 1
    for i, fila in enumerate(tab.index):
        rc = ws.cell(r0 + 3 + i, c0, str(fila)); rc.font = BOLD; rc.border = BORDER
        for j in range(len(tab.columns)):
            val = tab.iloc[i, j]
            cell = ws.cell(r0 + 3 + i, c0 + 1 + j)
            cell.border = BORDER; cell.alignment = Alignment(horizontal="center")
            if pd.isna(val):
                continue
            if modo == "riesgo":
                cell.value = round(float(val), 4)
                cell.fill = PatternFill("solid", fgColor=hex_color(norm_riesgo(val, vmin, vmax)))
            else:
                cell.value = int(val)
                g = int(255 - 120 * (val / mx if mx else 0))
                cell.fill = PatternFill("solid", fgColor=f"D6E4{g:02X}")
    ws.column_dimensions[get_column_letter(c0)].width = 26
    for j in range(len(tab.columns)):
        ws.column_dimensions[get_column_letter(c0 + 1 + j)].width = 16
    return r0 + 4 + len(tab.index)


def escribir_tabla_df(ws, r0, c0, titulo, df):
    if titulo:
        ws.cell(r0, c0, titulo).font = BOLD
        r0 += 1
    cols = list(df.columns)
    for j, h in enumerate(cols):
        cc = ws.cell(r0, c0 + j, str(h)); cc.font = BOLD; cc.border = BORDER
    for i in range(len(df)):
        for j, h in enumerate(cols):
            v = df.iloc[i][h]
            cell = ws.cell(r0 + 1 + i, c0 + j); cell.border = BORDER
            if isinstance(v, (int, float, np.integer, np.floating)) and not pd.isna(v):
                cell.value = round(float(v), 5)
            else:
                cell.value = str(v)
    return r0 + 2 + len(df)


def ejes_unicos(labels):
    """Garantiza labels únicos para pd.cut (agrega espacios si se repiten)."""
    vistos, out = {}, []
    for l in labels:
        if l in vistos:
            vistos[l] += 1
            out.append(l + " " * vistos[l])
        else:
            vistos[l] = 0
            out.append(l)
    return out


# ====================================================================================
# PROCESA UN ESCENARIO -> UNA HOJA
# ====================================================================================
def procesar_escenario(wb, nombre, df_sc, metodologia, excluir=()):
    ws = wb.create_sheet(title=nombre[:31])
    df_sc = df_sc.loc[pd.to_numeric(df_sc[RD_COL], errors="coerce").notna()].reset_index(drop=True)
    n = len(df_sc)
    if n < 50:
        ws.cell(1, 1, f"Escenario con muy pocos casos ({n}). Se omite.")
        return None, f"{'='*70}\nESCENARIO: {nombre}\n(omitido: solo {n} casos)\n\n"

    y = pd.to_numeric(df_sc[RD_COL], errors="coerce").values
    vmin, vmax = np.nanmin(y), np.nanmax(y)

    # ----- 1) Series crudas de predictores -----
    num = [c for c in NUM_VARS if c in df_sc.columns and c not in excluir]
    raw = {c: pd.to_numeric(df_sc[c], errors="coerce") for c in num}
    feats = list(num)
    if ORD_VAR in df_sc.columns:
        raw[ORD_VAR] = df_sc[ORD_VAR].map({f"G{i}": i for i in range(1, 9)})
        feats.append(ORD_VAR)
    if SCORE_COL in df_sc.columns and SIT_COL in df_sc.columns:
        bandas = banda_cortes_vec(pd.to_numeric(df_sc[SCORE_COL], errors="coerce").values,
                                  df_sc[SIT_COL].values)
        raw[SCORE_BANDA] = pd.Series(bandas, index=df_sc.index).map({f"G{i}": i for i in range(1, 6)})
        feats.append(SCORE_BANDA)

    # ----- 2) Dirección + (B) codificación WoE que entra al árbol -----
    usa_woe = (metodologia == "B" and WOE_AL_ARBOL)
    Xcols, misscol, cst_map, woemap, tablas, drop = {}, {}, {}, {}, [], []
    samp = (df_sc.sample(OPTB_SAMPLE, random_state=RANDOM_STATE).index
            if (metodologia == "B" and n > OPTB_SAMPLE) else df_sc.index)
    y_s = pd.Series(y, index=df_sc.index)
    for c in feats:
        serie = raw[c]
        if metodologia == "B" and c in NUM_VARS:                     # optbinning
            d = DIRECCION_NEGOCIO.get(c, 0)
            feasible, etiqueta, tb, dir_ef, optb = optbinning_fit(
                serie.loc[samp].values, y_s.loc[samp].values, c, d)
            tablas.append((c, etiqueta, tb))
            if not feasible:
                drop.append(c)
                continue
            if usa_woe:                                              # el WoE entra al árbol
                metric = "woe"
                try:
                    col = optb.transform(serie.values, metric="woe")
                    valcol = "WoE"
                except Exception:                                    # fallback: media de RD por bin
                    col = optb.transform(serie.values, metric="mean")
                    valcol = "Mean"
                Xcols[c] = np.asarray(col, dtype=float)
                cst_map[c] = +1                                      # WoE/mean ya crece con el RD
                woemap[c] = labelmap_from_table(tb, valcol)
                misscol[c] = np.zeros(n, dtype=bool)                 # missing ya codificado en el WoE
            else:
                Xcols[c] = serie.fillna(SENTINEL).values
                cst_map[c] = dir_ef
                misscol[c] = serie.isna().values
        else:                                                        # A (todas) / ordinales
            d = DIRECCION_NEGOCIO.get(c, 0)
            if d != 0:
                cst_map[c] = d
            else:
                ok = serie.notna()
                if ok.sum() < 30:
                    cst_map[c] = 0
                else:
                    rho, p = spearmanr(serie[ok], y[ok.values])
                    cst_map[c] = 0 if (np.isnan(rho) or p > 0.05) else int(np.sign(rho))
            Xcols[c] = serie.fillna(SENTINEL).values
            misscol[c] = serie.isna().values

    feats = [c for c in feats if c not in drop]
    X = pd.DataFrame({c: Xcols[c] for c in feats})
    miss = pd.DataFrame({c: misscol[c] for c in feats})
    cst = [cst_map[c] for c in feats]

    # ----- 3) Árbol monótono sobre RD -----
    try:
        tree = DecisionTreeRegressor(max_depth=MAX_DEPTH, min_samples_leaf=MIN_SAMPLES_LEAF,
                                     monotonic_cst=cst, random_state=RANDOM_STATE).fit(X, y)
    except (TypeError, ValueError):
        tree = DecisionTreeRegressor(max_depth=MAX_DEPTH, min_samples_leaf=MIN_SAMPLES_LEAF,
                                     random_state=RANDOM_STATE).fit(X, y)

    # ----- 4) Imagen del árbol (estructura; verde=RD bajo). Los RANGOS legibles van en las tablas. -----
    plt.figure(figsize=(26, 12))
    anns = plot_tree(tree, feature_names=feats, filled=True, rounded=True,
                     impurity=False, precision=4, fontsize=9, proportion=True)
    for ann, v in zip(anns, tree.tree_.value.reshape(-1)):
        if ann.get_bbox_patch() is not None:
            ann.get_bbox_patch().set_facecolor(CMAP(norm_riesgo(v, vmin, vmax)))
            ann.get_bbox_patch().set_edgecolor("black")
    plt.title("Árbol — verde = RD bajo, rojo = RD alto" +
              ("  (umbrales en WoE; ver 'cuadro del árbol completo' para los rangos)" if usa_woe else ""))
    plt.tight_layout(); plt.savefig(f"_tree_{nombre}.png", dpi=120, bbox_inches="tight"); plt.close()

    # ----- 5) Encabezado + imagen -----
    ws.cell(1, 1, f"ESCENARIO: {nombre}").font = Font(bold=True, size=14)
    ws.cell(2, 1, f"n = {n:,}   |   RD medio = {np.nanmean(y):.4f}   |   "
                  f"metodología = {'Solo árbol' if metodologia == 'A' else 'Optbinning + árbol (WoE al árbol)' if usa_woe else 'Optbinning + árbol'}")
    ws.cell(3, 1, f"monotonic_cst: {dict(zip(feats, cst))}")
    try:
        img = XLImage(f"_tree_{nombre}.png"); img.width, img.height = 1100, 520
        ws.add_image(img, "A5")
    except Exception as e:
        ws.cell(5, 1, f"(no se pudo insertar imagen: {e})")

    # ----- 6) Cuadros de doble entrada (top-2 variables) -----
    imp = pd.Series(tree.feature_importances_, index=feats).sort_values(ascending=False)
    top = [v for v in imp.index if imp[v] > 0][:2]
    for v in feats:
        if len(top) >= 2:
            break
        if v not in top:
            top.append(v)
    vfil, vcol = top[0], top[1]

    cuadro = {}
    for var in (vfil, vcol):
        idx = feats.index(var)
        es_woe = var in woemap
        thr = sorted({round(t, 6) for f, t in zip(tree.tree_.feature, tree.tree_.threshold)
                      if f == idx and (es_woe or t > SENTINEL / 2)})
        edges = [-np.inf] + thr + [np.inf]
        labels = ejes_unicos([cond_label(woemap, var, e0, e1) for e0, e1 in zip(edges[:-1], edges[1:])])
        b = pd.cut(X[var], bins=edges, labels=labels).astype(object)
        if not es_woe:
            b[miss[var].values] = "Missing"
        cuadro[var] = (b, labels)
    bf, lab_f = cuadro[vfil]
    bc, lab_c = cuadro[vcol]
    ord_f = lab_f + (["Missing"] if (bf == "Missing").any() else [])
    ord_c = lab_c + (["Missing"] if (bc == "Missing").any() else [])
    base = pd.DataFrame({"f": bf, "c": bc, "s": y})
    cant = base.pivot_table(index="f", columns="c", values="s", aggfunc="count").reindex(ord_f, axis=0).reindex(ord_c, axis=1)
    riesgo = base.pivot_table(index="f", columns="c", values="s", aggfunc="mean").reindex(ord_f, axis=0).reindex(ord_c, axis=1)

    fila = 34
    fila = escribir_cuadro(ws, fila, 1, "CUADRO POR CANTIDAD", cant, vfil, vcol, "cantidad") + 2
    fila = escribir_cuadro(ws, fila, 1, "CUADRO POR RIESGO (RD promedio)", riesgo,
                           vfil, vcol, "riesgo", vmin, vmax) + 2

    # ----- 6b) Importancia de variables -----
    ws.cell(fila, 1, "IMPORTANCIA DE VARIABLES (peso en el árbol)").font = Font(bold=True, size=12)
    for j, h in enumerate(["Variable", "Peso", "Peso %"]):
        c = ws.cell(fila + 1, 1 + j, h); c.font = BOLD; c.border = BORDER
    for i, (var, w) in enumerate(imp.items()):
        ws.cell(fila + 2 + i, 1, var).border = BORDER
        ws.cell(fila + 2 + i, 2, round(float(w), 4)).border = BORDER
        ws.cell(fila + 2 + i, 3, round(float(w) * 100, 1)).border = BORDER
    fila = fila + 3 + len(imp)

    # ----- 7) Cuadro del árbol completo (1 fila = 1 hoja, con LABELS de los bins) -----
    leafmap = reglas_por_hoja(tree, feats)
    usadas = [f for f in feats if any(f in cond for cond in leafmap.values())]
    leaf_id = tree.apply(X)
    try:
        quint_v = pd.qcut(y_s, 5, labels=["G1", "G2", "G3", "G4", "G5"])
    except ValueError:
        quint_v = pd.qcut(y_s.rank(method="first"), 5, labels=["G1", "G2", "G3", "G4", "G5"])
    res = pd.DataFrame({"leaf": leaf_id, "s": y, "quint": quint_v.values})

    filas = []
    for lid, sub in res.groupby("leaf"):
        cond = leafmap.get(lid, {})
        row = {f: cond_label(woemap, f, *cond.get(f, (-np.inf, np.inf))) for f in usadas}
        row["n"] = len(sub)
        row["%"] = round(100 * len(sub) / n, 1)
        row["RD_medio"] = sub["s"].mean()
        row["quintil_RD"] = sub["quint"].mode().iloc[0] if not sub["quint"].mode().empty else ""
        filas.append(row)
    tab = pd.DataFrame(filas).sort_values("RD_medio", ascending=True).reset_index(drop=True)

    ws.cell(fila, 1, "CUADRO DEL ÁRBOL COMPLETO (una fila = una hoja, ordenado por RD ascendente)").font = Font(bold=True, size=12)
    cols = usadas + ["n", "%", "RD_medio", "quintil_RD"]
    for j, h in enumerate(cols):
        c = ws.cell(fila + 1, 1 + j, h); c.font = BOLD; c.border = BORDER
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for i in range(len(tab)):
        for j, h in enumerate(cols):
            cell = ws.cell(fila + 2 + i, 1 + j); cell.border = BORDER
            cell.alignment = Alignment(horizontal="center")
            val = tab.iloc[i][h]
            if h == "RD_medio":
                cell.value = round(float(val), 4)
                cell.fill = PatternFill("solid", fgColor=hex_color(norm_riesgo(val, vmin, vmax)))
            elif h in ("n", "%"):
                cell.value = float(val) if h == "%" else int(val)
            else:
                cell.value = str(val)
    for j, h in enumerate(cols):
        ws.column_dimensions[get_column_letter(1 + j)].width = 20 if h in usadas else 11
    fila = fila + 3 + len(tab)

    # ----- 8) (Solo B) TABLA optbinning por variable (bin / WoE / RD medio) -----
    if tablas:
        fila += 2
        ws.cell(fila, 1, "TABLA OPTBINNING POR VARIABLE (bin · WoE · RD medio)").font = Font(bold=True, size=12)
        fila += 1
        for nom, etiqueta, tb in tablas:
            ws.cell(fila, 1, f"{nom}  —  {etiqueta}").font = BOLD
            if tb is not None:
                fila = escribir_tabla_df(ws, fila + 1, 1, "", tb) + 1
            else:
                fila += 2

    # ----- 9) Leads por APETITO de RD -----
    leaf_rd = res.groupby("leaf")["s"].mean().to_dict()
    regla_txt = {}
    for lid, cond in leafmap.items():
        partes = [f"{f} {cond_label(woemap, f, *rng)}" for f, rng in cond.items()
                  if cond_label(woemap, f, *rng) != "(todos)"]
        regla_txt[lid] = " & ".join(partes) if partes else "(raíz)"
    sid = df_sc[SUBJECT_COL].values if SUBJECT_COL in df_sc.columns else df_sc.index.values
    leads = pd.DataFrame({
        "subject_id": sid, "escenario": nombre,
        "leaf_key": [f"{nombre}#{l}" for l in leaf_id],
        "regla": [regla_txt.get(l, "") for l in leaf_id],
        "rd_cliente": y, "leaf_rd": [leaf_rd.get(l, np.nan) for l in leaf_id],
    })
    leads = leads[leads["leaf_rd"] <= APETITO_RD * max(FACTORES_RD)].reset_index(drop=True)

    # ----- 10) Espejo en TEXTO -----
    L = ["=" * 70, f"ESCENARIO: {nombre}",
         f"n = {n:,} | RD medio = {np.nanmean(y):.4f} | "
         f"metodología = {'Solo árbol' if metodologia == 'A' else 'Optbinning + árbol (WoE al árbol)' if usa_woe else 'Optbinning + árbol'}",
         f"monotonic_cst: {dict(zip(feats, cst))}",
         "-" * 70, "ÁRBOL (texto, rangos de los bins; value = RD):",
         arbol_texto(tree, feats, woemap),
         "-" * 70, "IMPORTANCIA DE VARIABLES (peso / peso %):",
         pd.DataFrame({"peso": imp.round(4), "peso_%": (imp * 100).round(1)}).to_string(),
         "-" * 70, f"CUADRO POR CANTIDAD (filas={vfil} / columnas={vcol}):",
         cant.to_string(na_rep=""),
         "-" * 70, f"CUADRO POR RIESGO - RD promedio (filas={vfil} / columnas={vcol}):",
         riesgo.round(4).to_string(na_rep=""),
         "-" * 70, "CUADRO DEL ÁRBOL COMPLETO (ordenado por RD ascendente):",
         tab.to_string(index=False)]
    if tablas:
        L += ["-" * 70, "TABLA OPTBINNING POR VARIABLE (bin · WoE · RD medio):"]
        for nom, etiqueta, tb in tablas:
            L += ["", f"### {nom} — {etiqueta}", tb.to_string() if tb is not None else "  (descartada)"]
    texto = "\n".join(L) + "\n\n"

    return leads, texto


# ====================================================================================
# ORQUESTACIÓN
# ====================================================================================
def construir_escenarios(df):
    out = []
    for v in VALORES_FLG:
        base = df[df[FLG_COL] == v]
        nombre = v.replace(">=", "ge").replace("=", "").replace("/", "_")
        out.append((nombre[:24], base))
        if FAR_COL in df.columns:
            far = base[pd.to_numeric(base[FAR_COL], errors="coerce") == 1]
            out.append((f"{nombre}_far1"[:28], far))
    return out


def _dedup(pool):
    return pool.sort_values("leaf_rd", ascending=True).drop_duplicates("subject_id", keep="first")


def hoja_estrategia(wb, leads_total, ruta):
    ws = wb.create_sheet("Estrategia")
    if leads_total.empty:
        ws.cell(1, 1, f"No hay segmentos con RD <= {APETITO_RD*max(FACTORES_RD):.4f}.")
        return f"{'='*70}\nHOJA DE ESTRATEGIA\n(sin segmentos bajo el apetito)\n"

    ded = _dedup(leads_total)
    seg = (ded.groupby("leaf_key")
           .agg(escenario=("escenario", "first"), regla=("regla", "first"),
                rd=("leaf_rd", "first"), n_leads=("subject_id", "size"))
           .reset_index().sort_values("rd", ascending=True).head(TOP_N_ESTRATEGIAS).reset_index(drop=True))
    seg.insert(0, "estrategia", range(1, len(seg) + 1))
    thr = sorted(APETITO_RD * f for f in FACTORES_RD)
    rdmin, rdmax = seg["rd"].min(), seg["rd"].max()

    ws.cell(1, 1, f"HOJA DE ESTRATEGIA — apetito RD base = {APETITO_RD:.4f} (un segmento entra si su RD <= apetito)").font = Font(bold=True, size=13)
    ws.cell(3, 1, "Leads únicos por estrategia según APETITO de RD:").font = BOLD
    ws.cell(4, 1, "Estrategia").font = BOLD; ws.cell(4, 1).border = BORDER
    ws.cell(4, 2, "RD segmento").font = BOLD; ws.cell(4, 2).border = BORDER
    for j, t in enumerate(thr):
        h = ws.cell(4, 3 + j, f"RD<={t:.4f}"); h.font = BOLD; h.border = BORDER
    for i, row in enumerate(seg.itertuples(index=False)):
        ws.cell(5 + i, 1, int(row.estrategia)).border = BORDER
        cr = ws.cell(5 + i, 2, round(float(row.rd), 4)); cr.border = BORDER
        cr.fill = PatternFill("solid", fgColor=hex_color(norm_riesgo(row.rd, rdmin, rdmax)))
        for j, t in enumerate(thr):
            ws.cell(5 + i, 3 + j, int(row.n_leads) if row.rd <= t else 0).border = BORDER
    tot_row = 5 + len(seg)
    ws.cell(tot_row, 1, "TOTAL").font = BOLD; ws.cell(tot_row, 1).border = BORDER
    for j, t in enumerate(thr):
        c = ws.cell(tot_row, 3 + j, int(seg.loc[seg["rd"] <= t, "n_leads"].sum())); c.font = BOLD; c.border = BORDER

    main_seg = seg[seg["rd"] <= APETITO_RD]
    main_leads = ded[ded["leaf_key"].isin(main_seg["leaf_key"])].copy()
    rank = dict(zip(seg["leaf_key"], seg["estrategia"]))
    main_leads["estrategia"] = main_leads["leaf_key"].map(rank)
    main_leads = main_leads.sort_values(["estrategia", "rd_cliente"])
    csv_path = ruta.replace(".xlsx", "_leads.csv")
    main_leads[["estrategia", "leaf_key", "escenario", "regla", "subject_id", "rd_cliente", "leaf_rd"]].to_csv(csv_path, index=False)

    r0 = tot_row + 3
    ws.cell(r0, 1, f"DETALLE corte principal (RD <= apetito {APETITO_RD:.4f})  |  "
                   f"leads únicos: {len(main_leads):,}  |  CSV: {csv_path}").font = Font(bold=True, size=12)
    headers = ["Estrategia", "Escenario", "Regla del segmento", "RD_segmento", "n_leads"]
    for j, h in enumerate(headers):
        c = ws.cell(r0 + 1, 1 + j, h); c.font = BOLD; c.border = BORDER
    for i, row in enumerate(main_seg.itertuples(index=False)):
        ws.cell(r0 + 2 + i, 1, int(row.estrategia)).border = BORDER
        ws.cell(r0 + 2 + i, 2, row.escenario).border = BORDER
        ws.cell(r0 + 2 + i, 3, row.regla).border = BORDER
        cs = ws.cell(r0 + 2 + i, 4, round(float(row.rd), 4)); cs.border = BORDER
        cs.fill = PatternFill("solid", fgColor=hex_color(norm_riesgo(row.rd, rdmin, rdmax)))
        ws.cell(r0 + 2 + i, 5, int(row.n_leads)).border = BORDER
    ws.column_dimensions["C"].width = 80
    for col in ("A", "B", "D", "E"):
        ws.column_dimensions[col].width = 16

    if len(main_leads) <= LEADS_EN_EXCEL_MAX:
        ws2 = wb.create_sheet("Leads")
        cols = ["estrategia", "escenario", "subject_id", "rd_cliente", "leaf_rd", "regla"]
        for j, h in enumerate(cols):
            ws2.cell(1, 1 + j, h).font = BOLD
        for i, row in enumerate(main_leads[cols].itertuples(index=False), start=2):
            for j, val in enumerate(row):
                ws2.cell(i, 1 + j, val if not isinstance(val, float) else round(val, 4))
    print(f"  Leads (apetito {APETITO_RD:.4f}) -> {csv_path}  ({len(main_leads):,} clientes únicos)")

    comp = seg.copy()
    for t in thr:
        comp[f"RD<={t:.4f}"] = np.where(comp["rd"] <= t, comp["n_leads"], 0)
    total = {f"RD<={t:.4f}": int(seg.loc[seg['rd'] <= t, 'n_leads'].sum()) for t in thr}
    T = ["=" * 70, f"HOJA DE ESTRATEGIA — apetito RD base = {APETITO_RD:.4f}",
         "-" * 70, "Escenarios de apetito (leads únicos por estrategia):",
         comp[["estrategia", "escenario", "rd", "n_leads"] + [f"RD<={t:.4f}" for t in thr]].to_string(index=False),
         "TOTAL por umbral: " + ", ".join(f"{k}={v:,}" for k, v in total.items()),
         "-" * 70, f"DETALLE corte principal (RD <= {APETITO_RD:.4f}) | leads únicos: {len(main_leads):,} | CSV: {csv_path}",
         main_seg.to_string(index=False)]
    return "\n".join(T) + "\n"


def construir_workbook(df, metodologia, ruta, excluir=()):
    wb = Workbook(); wb.remove(wb.active)
    idx = wb.create_sheet("Índice")
    metnom = "Solo árbol" if metodologia == "A" else ("Optbinning + árbol (WoE al árbol)" if WOE_AL_ARBOL else "Optbinning + árbol")
    idx.cell(1, 1, f"Desagregación de riesgo (RD) — {metnom}").font = Font(bold=True, size=14)
    idx.cell(2, 1, f"objetivo: {RD_COL} (RD, alto = MAYOR riesgo) | apetito base = {APETITO_RD:.4f}")
    if excluir:
        idx.cell(3, 1, f"Variables excluidas: {', '.join(excluir)}").font = Font(italic=True)
    r = 5
    leads_list, txt_blocks = [], []
    cab = (f"DESAGREGACIÓN DE RIESGO (RD) — {metnom}\n"
           f"objetivo: {RD_COL} (RD, alto = MAYOR riesgo) | apetito base = {APETITO_RD:.4f}\n"
           + (f"Variables excluidas: {', '.join(excluir)}\n" if excluir else ""))
    for nombre, sub in construir_escenarios(df):
        leads, texto = procesar_escenario(wb, nombre, sub, metodologia, excluir)
        txt_blocks.append(texto)
        if leads is not None and len(leads):
            leads_list.append(leads)
        idx.cell(r, 1, f"• {nombre}  (n={len(sub):,})"); r += 1

    leads_total = pd.concat(leads_list, ignore_index=True) if leads_list else pd.DataFrame()
    txt_estr = hoja_estrategia(wb, leads_total, ruta)
    wb.save(ruta)
    print(f"Generado: {ruta}")
    txt_path = ruta.replace(".xlsx", ".txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(cab + "\n" + "".join(txt_blocks) + "\n" + (txt_estr or ""))
    print(f"Generado: {txt_path}")


def main():
    df = pd.read_csv(DATA_PATH)
    df = df[~df[RD_COL].isna()]
    print(f"Base: {len(df):,} filas, {df.shape[1]} columnas")

    construir_workbook(df, "A", "arbol_solo_RD.xlsx")
    construir_workbook(df, "B", "optbinning_arbol_RD.xlsx")
    print("Listo: arbol_solo_RD.xlsx y optbinning_arbol_RD.xlsx")

    excluir = [c for c in VARS_EXCLUIR if c in df.columns]
    if excluir:
        construir_workbook(df, "A", "arbol_solo_RD_sinalgunasvariables.xlsx", excluir)
        construir_workbook(df, "B", "optbinning_arbol_RD_sinalgunasvariables.xlsx", excluir)
        print(f"Listo (sin {', '.join(excluir)}).")


if __name__ == "__main__":
    main()
