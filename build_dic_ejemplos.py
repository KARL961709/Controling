# -*- coding: utf-8 -*-
"""Excel didactico: como se calcula cada variable (castigo/saldos) con ejemplo."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = Workbook(); ws = wb.active; ws.title = "Como_se_calcula"
TEAL = PatternFill("solid", fgColor="007A72"); HEADW = Font(bold=True, color="FFFFFF")
GREY = PatternFill("solid", fgColor="EFEFEF"); BOLD = Font(bold=True)
YEL  = PatternFill("solid", fgColor="FFF3CD")
bd = Border(*[Side(style="thin", color="CCCCCC")]*4)
WRAP = Alignment(wrap_text=True, vertical="top")
CEN  = Alignment(horizontal="center", vertical="center")

r = 1
def title(txt, fill=TEAL, font=HEADW, span=6):
    global r
    c = ws.cell(r, 1, txt); c.font = font; c.fill = fill; c.alignment = Alignment(vertical="center")
    for j in range(1, span+1): ws.cell(r, j).fill = fill
    r += 1

def table(headers, rows, hfill=GREY):
    global r
    for j, h in enumerate(headers, 1):
        c = ws.cell(r, j, h); c.font = BOLD; c.fill = hfill; c.border = bd; c.alignment = CEN
    r += 1
    for row in rows:
        for j, v in enumerate(row, 1):
            c = ws.cell(r, j, v); c.border = bd; c.alignment = WRAP
        r += 1
    r += 1  # blank

ws.cell(r,1,"CÓMO SE CREA CADA VARIABLE — castigos, saldos, entidades, mora, meses").font = Font(bold=True, size=14)
r += 2

# ===== Ejemplo A =====
title("EJEMPLO A · RCC castigo — cliente JUAN (cuentas 8113/8123/8133, ventana 24m)")
table(["codmes","entidad","cuenta","condicion (días mora)","saldo"],
      [["202604","00002 (IBK)","8113","120","1,500"],
       ["202604","00003","8113","200","1,000"],
       ["202604","00003","8123","500","2,000"],
       ["202604","00007","8113","90","800"],
       ["202504","00003","8113","150","900"]])

# ===== Ejemplo B =====
title("EJEMPLO B · RCC castigo NO-IBK — cliente MARÍA (solo mes 202604, no-IBK)")
table(["codmes","entidad","cuenta","condicion (días)","saldo"],
      [["202604","00003","8113","300","1,200"],
       ["202604","00003","8123","900","4,000"],
       ["202604","00006","8113","600","2,000"]])

# ===== Ejemplo C =====
title("EJEMPLO C · t_360_cliente — un cliente, 6 meses (más reciente → más antiguo)")
table(["cod_mes","txs_prom","planilla_prom","tc_prom","pasivo_prom","pasivo_fdp"],
      [["202606","100","90","50","250","500"],
       ["202605","80","80","40","200","400"],
       ["202604","60","70","30","150","300"],
       ["202603","40","60","20","100","200"],
       ["202602","20","50","10","50","0"],
       ["202601","0","40","0","0","0"]])

# ===== Diccionario con cálculo y ejemplo =====
title("CÁLCULO DE CADA VARIABLE + RESULTADO EN EL EJEMPLO")
H = ["Variable","Qué mide","Cómo se calcula","Resultado en el ejemplo","Particularidad / nota"]
for j,h in enumerate(H,1):
    c = ws.cell(r,j,h); c.font=HEADW; c.fill=TEAL; c.border=bd; c.alignment=CEN
r += 1

D = [
 ["DEUDA_CAS","Saldo castigado asignado al cliente",
  "Cascada c→d→e→f: 0 si castigo IBK actual (c) o IBK U24M (d); s_cast_noibk si no-IBK actual (e); sld_u24 si no-IBK U24M (f); si no 0",
  "JUAN tiene IBK (rama c) → 0.  MARÍA solo no-IBK → s_cast_noibk = 7,200",
  "Gana la 1ra rama de la cascada (por eso JUAN sale 0 aunque tenga no-IBK)"],
 ["nro_entidades_castigo","Cuántas entidades te castigaron en el mes",
  "COUNT(DISTINCT entidad) en codmes=202604 (cuentas 81)",
  "JUAN → 3  (00002, 00003, 00007)",
  "Solo el mes actual (202604)"],
 ["monto_castigado_total","Saldo castigado total del mes",
  "SUM(saldo) en codmes=202604",
  "JUAN → 1,500+1,000+2,000+800 = 5,300",
  "Suma todas las entidades"],
 ["monto_castigado_ibk","Saldo castigado en IBK",
  "SUM(saldo) codmes=202604 y entidad='00002'",
  "JUAN → 1,500",""],
 ["monto_castigado_otros","Saldo castigado fuera de IBK",
  "SUM(saldo) codmes=202604 y entidad<>'00002'",
  "JUAN → 1,000+2,000+800 = 3,800",""],
 ["max_dias_mora_castigo","Peor mora (días) del castigo del mes",
  "MAX(condicion) en codmes=202604",
  "JUAN → 500",
  "Toma el MÁXIMO de días"],
 ["meses_desde_ultimo_castigo","Meses desde el ÚLTIMO mes con castigo",
  "date_diff('month', MAX(codmes), 202604)",
  "JUAN → 0  (último castigo = abr-2026)",
  "0 = castigo muy reciente"],
 ["meses_desde_primer_castigo","Meses desde el PRIMER castigo (ventana 24m)",
  "date_diff('month', MIN(codmes), 202604)",
  "JUAN → 12  (primer reporte = abr-2025)",
  "Mide antigüedad del castigo"],
 ["dm_cast_noibk","Días de mora del castigo NO-IBK (consolidado)",
  "max( MIN(condicion) por entidad ), mes 202604, no-IBK",
  "MARÍA → 00003 min(300,900)=300 ; 00006 min=600 → MAX(300,600)= 600",
  "⚠ MAX del MIN: conservador → SUBESTIMA la mora (vs max-max = 900)"],
 ["s_cast_noibk","Saldo castigado NO-IBK",
  "SUM(saldo) no-IBK en codmes=202604",
  "MARÍA → 1,200+4,000+2,000 = 7,200",
  "El saldo SÍ se suma completo (no se subestima)"],
 ["ANTI_CAST / FLG_CAST_AP","Antigüedad <5 / >=5 años y segmento",
  "dm_cast_noibk / 360 ; corte en 5 años (1800 días)",
  "MARÍA → 600/360 = 1.7 años → '<5años' → CAST_NOIBK_REP<5anios","60 meses = 5 años"],
 ["saldo_prom_tot_txs_um","Saldo txs último mes",
  "MAX del mes 202606",
  "= 100",
  "Patrón _um = mes 202606"],
 ["saldo_prom_tot_txs_u3m","Promedio txs U3M",
  "AVG(202606,202605,202604)",
  "(100+80+60)/3 = 80","_u3m = 3 meses"],
 ["saldo_prom_tot_txs_u6m","Promedio txs U6M",
  "AVG de los 6 meses",
  "(100+80+60+40+20+0)/6 = 50","_u6m = 6 meses. (planilla/tc siguen igual)"],
 ["saldo_prom_tot_pasivo_max_u6m","Máx pasivo en 6 meses",
  "MAX(pasivo_prom) en U6M",
  "MAX(250,200,150,100,50,0) = 250",""],
 ["var_pasivo_um_vs_u6m","Variación del pasivo",
  "(pasivo_um − pasivo_u6m) / pasivo_u6m",
  "um=250 ; u6m=(250+..+0)/6=125 → (250−125)/125 = 1.0","nullif evita /0 (null si u6m=0)"],
 ["saldo_pasivo_actual","Saldo pasivo (fdp) último mes",
  "MAX(pasivo_fdp) en 202606",
  "= 500",""],
 ["prom_saldo_pasivo_u4m","Promedio pasivo U4M",
  "AVG(202606,202605,202604,202603)",
  "(500+400+300+200)/4 = 350",""],
 ["max_saldo_pasivo_u6m","Máx pasivo (fdp) 6m",
  "MAX(pasivo_fdp)",
  "= 500",""],
 ["nro_meses_con_pasivo_u6m","Meses con pasivo > 0 (6m)",
  "COUNT(DISTINCT cod_mes con saldo>0)",
  "500,400,300,200>0 → 4  (los dos 0 no cuentan)",""],
 ["saldo_pasivo_componentes_u3m","Suma planilla+txs+tc (U3M)",
  "planilla_u3m + txs_u3m + tc_u3m",
  "80 + 80 + 40 = 200","Derivada (suma de 3 promedios)"],
 ["ratio_tc_pasivo_u3m","TC / pasivo (U3M)",
  "tc_u3m / nullif(pasivo_u3m,0)",
  "40 / 200 = 0.20","nullif → null si pasivo=0 (evita /0)"],
 ["nro_entidades_castigo_vida","# entidades castigo 'vida' (U24M)",
  "COUNT(DISTINCT entidad) [cuenta 81·(302/925), tipo_credito 11/12/13/99, codmes>=202403];  COALESCE(...,0)",
  "MARÍA → 2 ;  PEDRO sin castigo → 0",
  "★ COALESCE(...,0): el 0 = SIN castigo (de aquí salía el bin '<=0' del WoE)"],
 ["tipo_entidad_castigo_vida","Tipo de entidad del castigo vida",
  "'BIG FOUR' si entidad ∈ {00001,00002,00004,00006}, si no 'OTROS';  COALESCE(...,'SIN CASTIGO')",
  "MARÍA (tiene 00006) → BIG FOUR ;  PEDRO → SIN CASTIGO",
  "★ COALESCE(...,'SIN CASTIGO')"],
]
for row in D:
    for j,v in enumerate(row,1):
        c = ws.cell(r,j,v); c.border=bd; c.alignment=WRAP
        if "★ COALESCE" in str(row[4]): c.fill = YEL
    r += 1

ws.freeze_panes = "A2"
for col,w in zip("ABCDEF",[26,30,50,40,46,2]):
    ws.column_dimensions[col].width = w
wb.save("como_se_crea_variables.xlsx")
print("OK -> como_se_crea_variables.xlsx")
